"""Packet, sign-off pack and verdict logic of src/monthly_review.py (runs in a temporary outputs folder)."""
import json
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
import monthly_review as mr  # noqa: E402

GOOD = """# ABC Inc - Jun 2026
## Headline
1. Revenue was $195,898.16, down $12,452.80 (-6.0%) from May 2026.
## Revenue
2. Revenue is up $5,342.51 (+2.8%) from Jun 2025.
## Margins
3. Gross margin was 70.2%.
## Opex
4. Events & Conferences (+$17,970.67), Bad Debts (-$14,741.53) and Usage Fees (-$11,142.74) moved most.
## Cash & runway
5. Cash is $713,481.26; runway is about 3.4 months at the 3-month average burn.
## One-offs
6. The packet reports no one-off items in 2026-06.
## Questions
7. Is the Usage Fees amount expected to continue?
"""
FULL = {k: {"ok": 7, "of": 7} for k in ("numbers", "coverage", "causes", "tone")}
APPROVE = {"verdict": "approve", "changes": [], "statement_notes": [{"statement": i, "status": "ok", "note": ""} for i in range(1, 8)],
           "scores": FULL, "open_questions": [{"question": "Why did Usage Fees fall?", "why": "Largest revenue movement."}]}


class MonthlyReviewTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.out = Path(cls.tmp.name)
        cls.packet_path = mr.prepare("abc", "2026-06", outputs_root=cls.out)
        cls.packet = json.loads(cls.packet_path.read_text(encoding="utf-8"))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def build(self, commentary, review, month="2026-06"):
        d = self.out / "abc"
        (d / "c.md").write_text(commentary, encoding="utf-8")
        (d / "r.json").write_text(json.dumps(review), encoding="utf-8")
        return mr.build("abc", month, d / "c.md", d / "r.json", outputs_root=self.out)

    def test_gross_margin_drop_flags(self):
        self.assertEqual(self.packet["gross_margin_drop_threshold_pp"], 5)
        self.assertEqual([f["month"] for f in self.packet["gross_margin_flags"]], ["2025-03", "2025-08", "2026-04"])
        self.assertTrue(all(f["change_pp"] < -5 for f in self.packet["gross_margin_flags"]))

    def test_packet_has_source_cell_values(self):
        jun = self.packet["months"]["2026-06"]
        self.assertAlmostEqual(jun["revenue"]["value"], 195898.16, places=2)       # C14
        self.assertAlmostEqual(jun["net_income"]["value"], -330756.08, places=2)   # C107
        self.assertEqual([m["account"] for m in self.packet["top_movements"]], ["Events & Conferences", "Bad Debts", "Usage Fees"])
        self.assertAlmostEqual(self.packet["cash"]["value"], 713481.26, places=2)  # Balance Sheet D15

    def test_month_without_prior_is_rejected(self):
        with self.assertRaises(ValueError):
            mr.prepare("abc", "2025-01", outputs_root=self.out)

    def test_approved_pack_has_blank_signoff_and_open_questions(self):
        path, approved = self.build(GOOD, APPROVE)
        self.assertTrue(approved)
        wb = openpyxl.load_workbook(path)
        self.assertEqual(wb.sheetnames, ["Cover", "Summary", "Commentary Sign-off", "Numbers Check", "Top Movements", "Data checks", "Flux Items",
                                         "Reviewer Changes", "Open Questions"])
        sheet = wb["Commentary Sign-off"]
        self.assertEqual(sheet["G1"].value, "Sign-off (name / date)")
        self.assertTrue(all(row[6].value is None for row in sheet.iter_rows(min_row=2)))
        questions = [r[2] for r in wb["Open Questions"].iter_rows(min_row=2, values_only=True)]
        self.assertIn("Why did Usage Fees fall?", questions)
        self.assertTrue(any("Runway" in q for q in questions))  # auto question: runway under 6 months

    def test_wrong_number_blocks_approval_even_if_reviewer_approves(self):
        path, approved = self.build(GOOD.replace("70.2%", "71.2%"), APPROVE)
        self.assertFalse(approved)
        row = [r for r in openpyxl.load_workbook(path)["Commentary Sign-off"].iter_rows(min_row=2, values_only=True) if r[0] == 3][0]
        self.assertIn("MISMATCH", row[2])
        self.assertEqual(row[5], "CHANGE")

    def test_reviewer_changes_block_approval(self):
        review = {**APPROVE, "verdict": "changes", "changes": [{"statement": 1, "type": "invented_cause", "issue": "x", "fix": "y"}]}
        self.assertFalse(self.build(GOOD, review)[1])

    def test_missing_cash_statement_blocks_approval(self):
        self.assertFalse(self.build(GOOD.replace("5. Cash is $713,481.26; runway is about 3.4 months at the 3-month average burn.", "5. Cash is not discussed."), APPROVE)[1])

    def cover(self, path):
        return dict(openpyxl.load_workbook(path)["Cover"].iter_rows(min_row=2, values_only=True))

    def test_missing_section_blocks_approval_and_is_named(self):
        path, approved = self.build(GOOD.replace("## One-offs\n", ""), APPROVE)
        self.assertFalse(approved)
        self.assertIn("missing One-offs", self.cover(path)["Commentary sections"])

    def test_sections_out_of_order_are_not_accepted(self):
        swapped = GOOD.replace("## Margins", "## TMP").replace("## Opex", "## Margins").replace("## TMP", "## Opex")
        self.assertFalse(self.build(swapped, APPROVE)[1])

    def test_review_must_cite_statements_and_score_four_categories(self):
        no_scores = {k: v for k, v in APPROVE.items() if k != "scores"}
        self.assertFalse(self.build(GOOD, no_scores)[1])
        short = {**APPROVE, "scores": {**FULL, "tone": {"ok": 6, "of": 7}}}
        self.assertFalse(self.build(GOOD, short)[1])  # approving with a short score is inconsistent
        bad_cite = {**APPROVE, "verdict": "changes", "scores": {**FULL, "numbers": {"ok": 6, "of": 7}},
                    "changes": [{"type": "number_mismatch", "issue": "x", "fix": "y"}]}
        self.assertTrue(any("statement number" in p for p in mr.check_review(bad_cite, 7)))
        self.assertEqual(mr.check_review(APPROVE, 7), [])
        zero = {**bad_cite, "changes": [{"statement": 0, "type": "invented_cause", "issue": "x", "fix": "y"}]}
        self.assertTrue(any("statement 0" in p for p in mr.check_review(zero, 7)))

    def test_cover_data_checks_flux_and_memo(self):
        path, approved = self.build(GOOD, APPROVE)
        cover = self.cover(path)
        self.assertEqual(cover["Verdict"], "APPROVED")
        self.assertIn("Awaiting human approval", cover["Sign-off status"])
        self.assertEqual(cover["Net income excluding one-offs"], "-$330,756.08")
        self.assertIn("+2.8% vs 2025-06", cover["Revenue vs same month last year"])
        self.assertIn("$5,000 and 10%", cover["Flux threshold"])
        self.assertEqual(cover["Reviewer checklist"], "numbers 7/7, coverage 7/7, causes 7/7, tone 7/7")
        self.assertIn("graded PASS", cover["Owner memo"])
        wb = openpyxl.load_workbook(path)
        checks = list(wb["Data checks"].iter_rows(min_row=2, values_only=True))
        self.assertEqual({r[1] for r in checks}, {"PASS"})
        self.assertTrue(any("Balance Sheet balances" in r[0] for r in checks))
        flux = {r[0]: r for r in wb["Flux Items"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(flux["Events & Conferences"][8], "Income Statement!C65")
        self.assertIn("5,000", flux["Events & Conferences"][9])

    def test_owner_memo_is_graded_and_cash_comes_first(self):
        self.build(GOOD, APPROVE)
        memo = (self.out / "abc" / "owner_memo_2026-06.md").read_text(encoding="utf-8")
        self.assertLess(memo.index("## Cash and runway"), memo.index("## How the month went"))
        self.assertTrue(mr.grade_memo(memo, self.packet)["passed"])
        self.assertFalse(mr.grade_memo(memo.replace("$713,481", "$731,481"), self.packet)["passed"])
        self.assertIn("MoM", mr.grade_memo(memo + "MoM was bad.\n", self.packet)["jargon"])
        swapped = memo.replace("## Cash and runway", "## Other").replace("## How the month went", "## Cash and runway")
        self.assertFalse(mr.grade_memo(swapped, self.packet)["passed"])

    def test_memo_questions_lose_accountant_shorthand(self):
        self.assertEqual(mr.owner_memo.plain("TTM loss at the 3m burn, MoM?"), "twelve-month loss at the 3-month burn, month-on-month?")

    def test_yoy_matches_source_cells(self):
        y = self.packet["yoy"]
        self.assertEqual((y["month"], y["prior_year_month"]), ("2026-06", "2025-06"))
        self.assertAlmostEqual(y["revenue"]["value"], 195898.16, places=2)       # C14
        self.assertAlmostEqual(y["revenue"]["prior_value"], 190555.65, places=2)  # O14
        self.assertAlmostEqual(y["margins"]["gross_margin"]["prior_margin"], 141625.74 / 190555.65, places=5)  # O26 / O14
        self.assertEqual(list(y["opex_categories"]), ["Payroll & Benefits", "Sales & Marketing", "Technology"])

    def test_yoy_says_so_when_the_year_is_missing(self):
        mr.prepare("abc", "2025-02", outputs_root=self.out)
        p = json.loads((self.out / "abc" / "review_packet_2025-02.json").read_text(encoding="utf-8"))
        self.assertFalse(p["yoy"]["available"])
        self.assertIn("2024-02", p["yoy"]["reason"])

    def test_adjusted_net_income_removes_one_offs(self):
        ones = [{"account": "A", "section": "operating expenses", "month": "2025-12", "amount": 100.0},
                {"account": "B", "section": "income", "month": "2025-12", "amount": 40.0},
                {"account": "C", "section": "operating expenses", "month": "2025-11", "amount": 999.0}]
        a = mr.adjusted_net_income(-500.0, ones, "2025-12")
        self.assertEqual(a["one_off_effect_on_net_income"], -60.0)  # +40 income, -100 cost
        self.assertEqual(a["adjusted_net_income"], -440.0)
        self.assertEqual(self.packet["adjusted_net_income"]["adjusted_net_income"], self.packet["months"]["2026-06"]["net_income"]["value"])

    def test_materiality_comes_from_the_client_config(self):
        flux = json.loads((self.out / "abc" / "flux_2026-06.json").read_text(encoding="utf-8"))
        self.assertEqual(flux["materiality"]["source"], "client config")
        self.assertEqual((flux["materiality"]["min_dollars"], flux["materiality"]["min_pct"]), (5000.0, 0.1))
        ca = mr.close_artifacts
        loose = ca.build_flux(self.packet, ca.locate(mr.run_client.load_config("abc")), ca.materiality({"thresholds": {"materiality": {"min_dollars": 1000}}}))
        self.assertGreater(len(loose["items"]), len(flux["items"]))
        self.assertIn("$1,000", loose["materiality"]["statement"])
        with self.assertRaises(ValueError):
            ca.materiality({"thresholds": {"materiality": {"min_dolars": 1}}})

    def test_flux_needs_both_the_dollar_and_the_percent_threshold(self):
        ca = mr.close_artifacts
        li = lambda cur, pri: {"section": "operating expenses", "value": cur, "prior_value": pri,
                               "mom_change": round(cur - pri, 2), "mom_pct": round((cur - pri) / pri, 6)}
        packet = {"latest_month": "2026-06", "prior_month": "2026-05", "line_items": {
            "Big dollars, small percent": li(306000.0, 300000.0),   # +$6,000 but +2%: below min_pct
            "Big dollars, big percent": li(36000.0, 30000.0),       # +$6,000 and +20%: flagged
            "Small dollars, big percent": li(1400.0, 1000.0)}}      # +40% but only $400: below min_dollars
        flux = ca.build_flux(packet, {"accounts": {}, "month_cols": {}, "sheet": "x"}, ca.materiality())
        self.assertEqual([i["account"] for i in flux["items"]], ["Big dollars, big percent"])

    def test_commentary_numbering_is_enforced(self):
        with self.assertRaises(ValueError):
            mr.parse_statements("1. a\n3. b\n")

    def test_xyz_runs_without_cash_and_stays_in_its_folder(self):
        xyz = mr.prepare("xyz", "2026-03", outputs_root=self.out)
        self.assertEqual(xyz.parent.name, "xyz")
        self.assertNotIn("cash", json.loads(xyz.read_text(encoding="utf-8")))


if __name__ == "__main__":
    unittest.main()
