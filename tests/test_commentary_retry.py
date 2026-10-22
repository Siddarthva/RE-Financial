"""The /monthly-analysis grade-and-retry rule: on a failed grade the exact mismatches go back to the skill ONCE; a second
failure marks the pack NOT APPROVED with 'needs human' and the reasons. Never auto-approves."""
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
import commentary_retry as cr  # noqa: E402
import monthly_review  # noqa: E402

ABC = ROOT / "outputs" / "abc"
PACKET = json.loads((ABC / "review_packet_2026-06.json").read_text(encoding="utf-8"))
GOOD = (ABC / "commentary_2026-06.md").read_text(encoding="utf-8")
BAD = (ABC / "commentary_2026-06_flawed_control.md").read_text(encoding="utf-8")


class Reviser:
    """Fake skill: returns the queued texts and records the mismatches it was given."""
    def __init__(self, *texts):
        self.texts, self.calls = list(texts), []

    def __call__(self, text, mismatches):
        self.calls.append(mismatches)
        return self.texts.pop(0)


class RetryLogicTests(unittest.TestCase):
    def test_passes_first_time_without_retry(self):
        rev = Reviser()
        r = cr.grade_with_retry(GOOD, PACKET, rev)
        self.assertEqual((r["status"], r["retry_used"], r["needs_human"]), ("passed_first", False, False))
        self.assertEqual(rev.calls, [])

    def test_fails_then_passes_after_one_correction_with_exact_mismatches(self):
        rev = Reviser(GOOD)
        r = cr.grade_with_retry(BAD, PACKET, rev)
        self.assertEqual(r["status"], "passed_after_retry")
        self.assertTrue(r["retry_used"])
        self.assertEqual(len(rev.calls), 1)
        sent = rev.calls[0]
        self.assertEqual(sent, cr.mismatches(cr.grade(BAD, PACKET, require_sections=True)))  # the grader's own list, unchanged
        self.assertTrue(any(m.startswith("Top movement not mentioned: Usage Fees") for m in sent), sent)
        self.assertIn("Cash balance from the packet is not stated", sent)

    def test_fails_twice_is_not_approved_needs_human_and_never_a_third_try(self):
        rev = Reviser(BAD, GOOD)  # a third text is queued; it must never be requested
        r = cr.grade_with_retry(BAD, PACKET, rev)
        self.assertEqual((r["status"], r["needs_human"]), ("NOT APPROVED", True))
        self.assertEqual(len(rev.calls), 1)
        self.assertEqual(len(r["attempts"]), 2)
        self.assertTrue(r["reasons"])


class RetryCliAndPackTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.out = Path(self.tmp.name) / "outputs"
        d = self.out / "abc"
        d.mkdir(parents=True)
        for name in ("review_packet_2026-06.json", "intake_2026-06.json", "flux_2026-06.json", "review_2026-06.json"):
            shutil.copy2(ABC / name, d / name)
        (d / "commentary_2026-06.md").write_text(BAD, encoding="utf-8")

    def tearDown(self):
        self.tmp.cleanup()

    def cli(self, action):
        return cr.main([action, "abc", "2026-06", "--outputs-root", str(self.out)])

    def record(self):
        return json.loads((self.out / "abc" / "commentary_grade_2026-06.json").read_text(encoding="utf-8"))

    def test_cli_flow_and_pack_marked_not_approved(self):
        self.assertEqual(self.cli("regrade"), 2)  # no failed first grade yet
        self.assertEqual(self.cli("grade"), 1)
        self.assertEqual(self.record()["status"], "retry_pending")
        self.assertEqual(self.cli("regrade"), 1)  # the "correction" did not fix it
        rec = self.record()
        self.assertEqual((rec["status"], rec["needs_human"], rec["retry_used"]), ("NOT APPROVED", True, True))
        self.assertEqual(self.cli("regrade"), 2)  # no third attempt
        d = self.out / "abc"
        monthly_review.build("abc", "2026-06", d / "commentary_2026-06.md", d / "review_2026-06.json",
                             monthly_review.run_client.CLIENTS_ROOT, self.out)
        wb = load_workbook(d / "monthly_review_2026-06.xlsx", read_only=True)
        cover = {r[0]: r[1] for r in wb["Cover"].iter_rows(values_only=True)}
        wb.close()
        self.assertTrue(str(cover["Verdict"]).startswith("NOT APPROVED - needs human: "), cover["Verdict"])
        self.assertEqual(cover["Commentary grade"], "NOT APPROVED (retry used: yes)")
        self.assertFalse((d / "approval_2026-06.json").exists())

    def test_good_commentary_passes_first_and_pack_is_not_forced_down(self):
        (self.out / "abc" / "commentary_2026-06.md").write_text(GOOD, encoding="utf-8")
        self.assertEqual(self.cli("grade"), 0)
        self.assertEqual(self.record()["status"], "passed_first")


if __name__ == "__main__":
    unittest.main()
