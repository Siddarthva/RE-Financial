"""Layout of every generated workbook (analysis.xlsx for abc/xyz, the abc reviewer pack): no blank rows/columns
inside the used range, header in row 1 (or 3 under a title), frozen panes, filters, number formats, no empty sheets."""
import json
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import excel_layout  # noqa: E402
import monthly_review as mr  # noqa: E402
from run_client import run  # noqa: E402
from test_monthly_review import APPROVE, GOOD  # noqa: E402


def _filled(c):
    return c.value not in (None, "")


class ExcelLayoutTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        out = Path(cls.tmp.name)
        cls.books = {"xyz/analysis.xlsx": run("xyz", outputs_root=out)["workbook"]}
        mr.prepare("abc", "2026-06", outputs_root=out)  # also writes abc/analysis.xlsx
        cls.books["abc/analysis.xlsx"] = out / "abc" / "analysis.xlsx"
        (out / "abc" / "c.md").write_text(GOOD, encoding="utf-8")
        (out / "abc" / "r.json").write_text(json.dumps(APPROVE), encoding="utf-8")
        cls.books["abc/monthly_review"] = mr.build("abc", "2026-06", out / "abc" / "c.md", out / "abc" / "r.json", outputs_root=out)[0]
        cls.wbs = {k: openpyxl.load_workbook(p) for k, p in cls.books.items()}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def sheets(self):
        for name, wb in self.wbs.items():
            for ws in wb.worksheets:
                with self.subTest(book=name, sheet=ws.title):
                    yield ws

    def test_no_empty_sheets(self):
        for ws in self.sheets():
            self.assertTrue(any(_filled(c) for row in ws.iter_rows() for c in row))

    def test_header_in_row_1_or_3(self):
        for ws in self.sheets():
            first = min(c.row for row in ws.iter_rows() for c in row if _filled(c))
            self.assertIn(first, (1, 3))
            self.assertTrue(all(_filled(c) for c in ws[first]), "every column has a header")

    def test_no_blank_rows_or_columns_inside_used_range(self):
        for ws in self.sheets():
            cells = [c for row in ws.iter_rows() for c in row if _filled(c)]
            rows, cols = {c.row for c in cells}, {c.column for c in cells}
            self.assertEqual(sorted(rows), list(range(min(rows), max(rows) + 1)), "blank row inside data")
            self.assertEqual(sorted(cols), list(range(1, max(cols) + 1)), "blank column inside data")

    def test_frozen_header_and_first_column(self):
        for ws in self.sheets():
            self.assertIn(ws.freeze_panes, ("B2", "B4"))

    def test_print_setup_and_widths(self):
        for ws in self.sheets():
            self.assertEqual(ws.page_setup.orientation, "landscape")
            self.assertEqual(ws.page_setup.fitToWidth, 1)
            for col in range(1, ws.max_column + 1):
                width = ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width
                self.assertTrue(excel_layout.MIN_WIDTH <= width <= excel_layout.MAX_WIDTH, f"col {col} width {width}")

    def test_tables_have_filters(self):
        for ws in self.sheets():
            if ws.title not in ("Cover", "Summary"):  # key-value sheets, no table
                self.assertEqual(ws.auto_filter.ref, ws.dimensions)

    def test_number_and_date_formats(self):
        trends = self.wbs["abc/analysis.xlsx"]["Monthly Trends"]
        head = [c.value for c in trends[1]]
        row = trends[2]
        self.assertIsInstance(row[head.index("month")].value, datetime)
        self.assertEqual(row[head.index("month")].number_format, excel_layout.DATE)
        self.assertEqual(row[head.index("revenue")].number_format, excel_layout.MONEY)
        self.assertEqual(row[head.index("gross_margin_pct")].number_format, excel_layout.PCT)
        self.assertIn("[Red](", excel_layout.MONEY)

    def test_wrap_only_in_note_columns(self):
        for ws in self.sheets():
            for c in ws[1]:
                wrapped = any(cell.alignment.wrap_text for (cell,) in ws.iter_rows(min_row=2, min_col=c.column, max_col=c.column))
                if wrapped:
                    self.assertIn(c.value, mr.NOTE_COLUMNS | {"warning", "commentary"})


if __name__ == "__main__":
    unittest.main()
