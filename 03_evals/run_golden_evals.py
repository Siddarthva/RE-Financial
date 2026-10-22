"""Run deterministic golden cases for the monthly financial-analysis engine."""
from __future__ import annotations
import json
import math
import re
import sys
from pathlib import Path
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "03_evals"))
from financial_analysis import normalize_monthly, category_summary, make_commentary  # noqa: E402
import abc_workbook_adapter as abc  # noqa: E402

CASES = ROOT / "03_evals" / "golden_cases.json"
WORKBOOK = ROOT / "data" / "ABC Inc - Financial Statements.xlsx"
CENT = 0.005
_abc = {}


def abc_outputs():
    """Adapter outputs for the real workbook, read-only, computed once per run."""
    if not _abc:
        am = abc.parse_wide_income_statement(WORKBOOK)
        _abc.update(account_month=am, monthly=abc.build_monthly_summary(am).set_index("month"),
                    categories=abc.build_expense_category_trends(am),
                    balance_sheet=abc.parse_wide_income_statement(WORKBOOK, abc.ABC_BS_LAYOUT))
    return _abc


_cells = {}


def source_cells():
    """Cached Income Statement cell values, read-only and in memory: {(row, col letter): value}."""
    if not _cells:
        ws = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)["Income Statement"]
        for r in ws.iter_rows(min_row=1, max_row=110, max_col=20):
            for c in r:
                if c.value is not None:
                    _cells[(c.row, c.column_letter)] = c.value
    return _cells


def perturbed_cells(case):
    """Copy of the source cells with the case's perturbation applied; refuses if a recorded original is wrong."""
    cells = dict(source_cells())
    for ch in case["perturbation"]["cells_changed"]:
        m = re.fullmatch(r"([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?", ch["cell"])
        c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3) or m.group(1), int(m.group(4) or m.group(2))
        if c1 == c2 and r1 == r2:
            if cells.get((r1, c1)) != ch["original"]:
                raise ValueError(f"{ch['cell']}: recorded original {ch['original']!r} != workbook {cells.get((r1, c1))!r}")
        for col in [chr(x) for x in range(ord(c1), ord(c2) + 1)]:
            for row in range(r1, r2 + 1):
                if ch["new"] is None:
                    cells.pop((row, col), None)
                else:
                    cells[(row, col)] = ch["new"]
    return cells


def real_frame(case):
    """Month/value frame for the engine from perturbed workbook cells."""
    cells, inp = perturbed_cells(case), case["input"]
    cols = inp["columns"]
    if cols == "all":
        cols = [chr(x) for x in range(ord("C"), ord("T") + 1) if (5, chr(x)) in cells]
    df = pd.DataFrame({"month": [cells[(5, c)] for c in cols]})
    df["month"] = pd.to_datetime(df["month"].astype(str), format="mixed")
    def label(row):
        return cells.get((row, "A")) or cells.get((row, "B"))
    if "revenue_row" in inp and label(inp["revenue_row"]) == "Total Income":
        df["revenue"] = [cells[(inp["revenue_row"], c)] for c in cols]
    if "gross_profit_row" in inp:
        df["gross_profit"] = [cells[(inp["gross_profit_row"], c)] for c in cols]
    return df, cells, cols


def close(actual, expected, tol=CENT):
    return actual is not None and not pd.isna(actual) and abs(float(actual) - float(expected)) <= tol


def run_abc_case(typ, data, expected):
    o = abc_outputs()
    am, monthly, cats = o["account_month"], o["monthly"], o["categories"]
    if typ == "abc_seasonal_peak":
        errors = []
        for account, by_month in expected["values"].items():
            row = am.loc[account]
            for year, month in expected["peak_month_by_year"].items():
                year_row = row[pd.DatetimeIndex(row.index).year == int(year)]
                if year_row.idxmax() != pd.Timestamp(month):
                    errors.append(f"{account} {year} peak={year_row.idxmax():%Y-%m}")
            errors += [f"{account} {m}={row[pd.Timestamp(m)]}" for m, v in by_month.items() if not close(row[pd.Timestamp(m)], v)]
        return not errors, "; ".join(errors) or "February peak in 2025 and 2026"
    if typ == "abc_margin_flags":
        got = {f["month"]: f for f in abc.margin_drop_flags(monthly, data["threshold_pp"])}
        want = {f["month"]: f for f in expected["flags"]}
        if set(got) != set(want):
            return False, f"flagged {sorted(got)} expected {sorted(want)}"
        bad = [m for m in want if not close(got[m]["gross_margin_pct"], want[m]["gross_margin"], 1e-6)
               or not close(got[m]["change_pp"], want[m]["change_pp"], 1e-4)]
        return not bad, "; ".join(bad) or f"{len(want)} months flagged: {', '.join(sorted(want))}"
    if typ == "abc_one_off":
        found = abc.find_one_off_items(am, data["min_abs"])
        got = {(r.account, r.month.strftime("%Y-%m-%d")): r.amount for r in found.itertuples()}
        want = {(i["account"], i["month"]): i["amount"] for i in expected["one_off_items"]}
        ok = got.keys() == want.keys() and all(close(got[k], want[k]) for k in want)
        event = pd.Timestamp(expected["event_month"])
        positive_ni = list(monthly.index[monthly["net_income"] > 0])
        ok = ok and positive_ni == [event]
        ok = ok and all(close(monthly.at[event, k], v) for k, v in expected["event_month_totals"].items())
        # Removing an OPEX item from the month moves net income by +amount (expenses reduce NI).
        adjusted = monthly.at[event, "net_income"] + sum(
            v for (a, m), v in got.items() if pd.Timestamp(m) == event and found.set_index("account").at[a, "section"] == "operating expenses")
        ok = ok and close(adjusted, expected["net_income_excluding_one_offs"])
        return ok, f"one-offs={sorted(got)}; positive NI months={[f'{m:%Y-%m}' for m in positive_ni]}; NI ex one-offs={adjusted:,.2f}"
    if typ == "abc_first_ttm":
        first = pd.Timestamp(expected["first_valid_month"])
        valid = monthly.index[monthly["revenue_ttm"].notna()]
        ok = len(valid) > 0 and valid[0] == first and pd.isna(monthly.at[pd.Timestamp(expected["prior_month_ttm_is_null"]), "revenue_ttm"])
        # Every month before the first valid one must be unavailable for every TTM metric, not just the month before.
        earlier = monthly.loc[monthly.index < first, [f"{k}_ttm" for k in expected["ttm"]]]
        ok = ok and earlier.isna().all().all()
        ok = ok and all(close(monthly.at[first, f"{k}_ttm"], v) for k, v in expected["ttm"].items())
        return ok, f"first valid TTM={valid[0]:%Y-%m} revenue_ttm={monthly.at[first, 'revenue_ttm']:,.2f}; {len(earlier)} earlier months all null={earlier.isna().all().all()}"
    if typ == "abc_line_item_step":
        row, month = am.loc[data["account"]], pd.Timestamp(data["month"])
        prior = row[month - pd.DateOffset(months=1)]
        held = row[(row.index >= month) & (row.index <= pd.Timestamp(expected["held_through"]))]
        ok = (close(prior, expected["prior_value"]) and close(row[month], expected["value"])
              and close(row[month] - prior, expected["mom_change"]) and close((row[month] - prior) / abs(prior), expected["mom_pct"], 1e-6)
              and all(close(v, expected["value"]) for v in held))
        return ok, f"{prior:,.2f} -> {row[month]:,.2f} ({row[month] - prior:+,.2f}); held {len(held)} months"
    if typ == "abc_line_item_value":
        row = am.loc[data["account"]]
        errors = [f"{m}={row[pd.Timestamp(m)]}" for m, v in expected["values"].items() if not close(row[pd.Timestamp(m)], v)]
        may = cats[(cats["expense_category"] == data["account"]) & (cats["month"] == pd.Timestamp("2026-05-01"))].iloc[0]
        if expected["mom_pct_may_2026"] is None and not pd.isna(may["mom_pct"]):
            errors.append(f"May 2026 mom_pct={may['mom_pct']} expected null")
        return not errors, "; ".join(errors) or f"values match; May 2026 mom_pct={may['mom_pct']}"
    if typ == "abc_cash_runway":
        pos = abc.cash_position(o["balance_sheet"], data["cash_row"], windows=tuple(int(n) for n in expected["burn"]))
        ok = (pos["month"] == pd.Timestamp(expected["month"]) and close(pos["cash"], expected["cash"])
              and close(pos["prior_cash"], expected["prior_cash"])
              and all(close(pos["burn"][int(n)], v) for n, v in expected["burn"].items())
              and all(close(pos["runway_months"][int(n)], v, 1e-6) for n, v in expected["runway_months"].items()))
        primary = int(data["primary_window_months"])
        return ok, f"cash={pos['cash']:,.2f}; {primary}m burn={pos['burn'][primary]:,.2f}; runway={pos['runway_months'][primary]:.2f} months"
    if typ == "abc_opex_spike":
        one_offs = set(abc.find_one_off_items(am, data["exclude_one_offs_min_abs"])["account"])
        details = []
        # A spike is an increase to a positive expense; a credit reversing to zero is not a spike.
        spikes = cats[cats["amount"] > 0]
        for key, pool in [("largest_overall", spikes), ("largest_recurring", spikes[~spikes["expense_category"].isin(one_offs)])]:
            top = pool.loc[pool["mom_change"].idxmax()]
            exp = expected[key]
            ok = (top["expense_category"] == exp["account"] and top["month"] == pd.Timestamp(exp["month"])
                  and close(top["amount"], exp["amount"]) and close(top["mom_change"], exp["mom_change"])
                  and (pd.isna(top["mom_pct"]) if exp["mom_pct"] is None else close(top["mom_pct"], exp["mom_pct"], 1e-6)))
            details.append((ok, f"{key}={top['expense_category']} {top['month']:%Y-%m} +{top['mom_change']:,.2f} pct={top['mom_pct']}"))
        return all(d[0] for d in details), "; ".join(d[1] for d in details)
    if typ == "abc_latest_summary":
        month = pd.Timestamp(data["month"])
        errors = []
        for metric, fields in expected.items():
            for field, v in fields.items():
                col = {"value": metric, "ttm": f"{metric}_ttm", "mom_change": f"{metric}_mom_change",
                       "mom_pct": f"{metric}_mom_pct", "margin": {"gross_profit": "gross_margin_pct",
                       "operating_income": "operating_margin_pct", "net_income": "net_margin_pct"}.get(metric)}[field]
                tol = 1e-6 if field in {"mom_pct", "margin"} else CENT
                actual = monthly.at[month, col] if col in monthly.columns else None
                if not close(actual, v, tol):
                    errors.append(f"{metric}.{field}={actual} expected {v}")
        return not errors, "; ".join(errors) or "all summary fields match"
    return False, f"Unknown case type: {typ}"


def approx(actual, expected, tol=1e-6):
    return actual is not None and not pd.isna(actual) and math.isclose(float(actual), float(expected), abs_tol=tol, rel_tol=tol)


def run_real_case(case):
    """G01-G08 and G14: the same engine rules, fed from (perturbed) real workbook cells."""
    typ, expected = case["type"], case["expected"]
    df, cells, cols = real_frame(case)
    if typ in {"revenue_mom", "zero_prior_revenue"}:
        out = normalize_monthly(df)
        change, pct = out.iloc[-1]["revenue_mom_change"], out.iloc[-1]["revenue_mom_pct"]
        ok = close(change, expected["absolute_change"])
        ok = ok and (pd.isna(pct) if expected["percentage_change"] is None else approx(pct, expected["percentage_change"]))
        return ok, f"change={change}, pct={pct}"
    if typ == "month_labels":
        out = normalize_monthly(df)
        got = [f"{m:%Y-%m-%d}" for m in out["month"]]
        ok = got == expected["months"] and all(close(a, b) for a, b in zip(out["revenue"], expected["revenue"]))
        return ok, f"months={got}, revenue={out['revenue'].tolist()}"
    if typ == "gross_margin":
        actual = normalize_monthly(df).iloc[0]["gross_margin_pct"]
        return approx(actual, expected["gross_margin"]), f"gross_margin={actual}"
    if typ == "ttm":
        out = normalize_monthly(df)
        before, after = out.iloc[10]["revenue_ttm"], out.iloc[11]["revenue_ttm"]
        return pd.isna(before) and close(after, expected["ttm_after_12"]), f"ttm_after_11={before}, ttm_after_12={after}"
    if typ in {"duplicate_month", "missing_revenue"}:
        try:
            normalize_monthly(df)
        except ValueError as exc:
            return expected["raises"] in str(exc), str(exc)
        return False, "No validation error was raised"
    if typ == "expense_category":
        amounts = [cells[(case["input"]["category_row"], c)] for c in cols]
        long = pd.DataFrame({"month": df["month"], "expense_category": expected["category"], "expense_amount": amounts})
        last = category_summary(long).iloc[-1]
        ok = close(last["expense_mom_change"], expected["absolute_change"]) and approx(last["expense_mom_pct"], expected["percentage_change"])
        return ok, f"change={last['expense_mom_change']}, pct={last['expense_mom_pct']}"
    if typ == "short_history_commentary":
        out = normalize_monthly(df)
        text = " ".join(make_commentary(out))
        return len(out) == expected["months_remaining"] and expected["commentary_contains"].lower() in text.lower(), f"{len(out)} months; {text}"
    return False, f"Unknown case type: {typ}"


def run_intake_case(case):
    """G19-G22: a perturbed values-only copy of the workbook (temp file) must be refused by the intake gate
    (close_artifacts.intake_gate, shared by run_pipeline and /monthly-analysis) with a message naming the cause."""
    import io
    import tempfile
    import close_artifacts
    import run_client
    pert = case["perturbation"]
    if pert.get("file") == "formulas_without_cached_values":
        wb = openpyxl.load_workbook(io.BytesIO(WORKBOOK.read_bytes()), data_only=False)
    else:
        wb = openpyxl.load_workbook(io.BytesIO(WORKBOOK.read_bytes()), data_only=True)
        for ch in pert.get("cells_changed", []):
            cell = wb[ch["sheet"]][ch["cell"]]
            if cell.value != ch["original"]:
                raise ValueError(f"{ch['cell']}: recorded original {ch['original']!r} != workbook {cell.value!r}")
            cell.value = ch["new"]
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "perturbed.xlsx"
        wb.save(path)
        cfg = {**run_client.load_config("abc"), "workbook_path": path}
        try:
            report, _, _ = close_artifacts.intake_gate(cfg, case["input"]["month"], abc)
            message = close_artifacts.refusal(report, case["input"]["month"])
        except ValueError as exc:
            message = str(exc)
    if not message:
        return False, "intake gate accepted the perturbed workbook"
    missing = [s for s in case["expected"]["message_contains"] if s not in message]
    return not missing, (f"missing {missing} in: " if missing else "refused: ") + message[:160]


def run_flux_item_case(case):
    """G26-G28: the flux table's row for one account equals values read by hand from the source cells."""
    import close_artifacts
    import monthly_review
    import run_client
    inp, exp = case["input"], case["expected"]
    cells = source_cells()
    for ref, want in case["source_cells"].items():
        m = re.fullmatch(r"([A-Z]+)(\d+)", ref)
        if not close(cells.get((int(m.group(2)), m.group(1))), want):
            return False, f"source cell {ref} is {cells.get((int(m.group(2)), m.group(1)))!r}, case records {want!r}"
    a = abc_outputs()
    month = pd.Timestamp(inp["month"] + "-01")
    packet = monthly_review.build_computed(a["account_month"], a["monthly"], month, None)
    flux = close_artifacts.build_flux(packet, close_artifacts.locate(run_client.load_config("abc")))
    row = next((i for i in flux["items"] if i["account"] == inp["account"]), None)
    if row is None:
        return False, f"{inp['account']} not flagged"
    ok = row["reason_code"] == exp["reason_code"] and row["source_cell_current"] == exp["source_cell_current"]         and row["source_cell_prior"] == exp["source_cell_prior"]         and all(close(row[k], exp[k]) for k in ("current", "prior", "dollar_change"))
    return ok, f"{row['reason_code']} {row['current']} vs {row['prior']} ({row['dollar_change']:+,.2f}) at {row['source_cell_current']}"


def run_skill_text_case(case):
    """G29-G30: the saved skill output for ABC 2026-06 passes its grader against the engine JSON."""
    import grade_workflow_skills as gw
    inp = case["input"]
    d = ROOT / "outputs" / inp["client"]
    flux = json.loads((d / f"flux_{inp['month']}.json").read_text(encoding="utf-8"))
    text = (d / inp["file"].format(month=inp["month"])).read_text(encoding="utf-8")
    if inp["grader"] == "flux":
        problems = gw.grade_flux(text, flux)
    else:
        review = json.loads((d / f"review_{inp['month']}.json").read_text(encoding="utf-8"))
        packet = json.loads((d / f"review_packet_{inp['month']}.json").read_text(encoding="utf-8"))
        problems = gw.grade_questions(text, flux, review, packet)
    return not problems, "grader passed" if not problems else "; ".join(problems[:3])


def _check_source_cells(case):
    cells = source_cells()
    for ref, want in case["source_cells"].items():
        m = re.fullmatch(r"([A-Z]+)(\d+)", ref)
        if not close(cells.get((int(m.group(2)), m.group(1))), want):
            return f"source cell {ref} is {cells.get((int(m.group(2)), m.group(1)))!r}, case records {want!r}"
    return None


def run_yoy_case(case):
    """G31: the packet's year-over-year block equals values read by hand from the source cells."""
    import monthly_review
    bad = _check_source_cells(case)
    if bad:
        return False, bad
    a = abc_outputs()
    y = monthly_review.build_computed(a["account_month"], a["monthly"], pd.Timestamp(case["input"]["month"] + "-01"), None)["yoy"]
    e, g = case["expected"], y["margins"]["gross_margin"]
    ok = (y["prior_year_month"] == e["prior_year_month"] and close(y["revenue"]["value"], e["revenue"]) and close(y["revenue"]["prior_value"], e["prior_revenue"])
          and close(y["revenue"]["yoy_change"], e["yoy_change"]) and close(y["revenue"]["yoy_pct"], e["yoy_pct"], 1e-5)
          and close(g["margin"], e["gross_margin"], 1e-5) and close(g["prior_margin"], e["prior_gross_margin"], 1e-5)
          and close(g["yoy_change_pp"], e["yoy_change_pp"], 1e-4))
    return ok, f"revenue {y['revenue']['value']} vs {y['revenue']['prior_value']} ({y['revenue']['yoy_pct']:+.4%}); gross margin change {g['yoy_change_pp']:+.4f} pp"


def run_adjusted_case(case):
    """G32: net income excluding the month's one-offs equals the hand calculation from the source cells."""
    import monthly_review
    bad = _check_source_cells(case)
    if bad:
        return False, bad
    a = abc_outputs()
    month = case["input"]["month"]
    ones = [{"account": r.account, "section": r.section, "month": f"{r.month:%Y-%m}", "amount": r.amount}
            for r in abc.find_one_off_items(a["account_month"]).itertuples()]
    ni = a["monthly"].at[pd.Timestamp(month + "-01"), "net_income"]
    got, e = monthly_review.adjusted_net_income(ni, ones, month), case["expected"]
    ok = (close(got["net_income"], e["net_income"]) and close(got["one_off_effect_on_net_income"], e["one_off_effect_on_net_income"])
          and close(got["adjusted_net_income"], e["adjusted_net_income"]) and sorted(got["items"]) == sorted(e["items"]))
    return ok, f"net income {got['net_income']:,.2f}, one-off effect {got['one_off_effect_on_net_income']:+,.2f}, adjusted {got['adjusted_net_income']:,.2f}"


def run_sections_case(case):
    """G33: the saved ABC commentary has every required section, in order."""
    import grade_commentary
    inp = case["input"]
    text = (ROOT / "outputs" / inp["client"] / inp["file"].format(month=inp["month"])).read_text(encoding="utf-8")
    r = grade_commentary.check_sections(text)
    return r["passed"] == case["expected"]["passed"] and r["found"] == case["expected"]["sections"], f"found {r['found']}, missing {r['missing']}"


_packets = {}


def _report_field(f):
    """Value of one named field in a saved QuickBooks report JSON: the row whose first cell is f['row'], column f['column']."""
    rep = json.loads((ROOT / f["file"]).read_text(encoding="utf-8"))
    titles = [c.get("ColTitle") for c in rep["Columns"]["Column"]]

    def walk(rows):
        for x in rows.get("Row", []):
            for part in ("Header", "Summary"):
                if part in x:
                    yield x[part]["ColData"]
            if "ColData" in x:
                yield x["ColData"]
            if "Rows" in x:
                yield from walk(x["Rows"])
    cd = next(c for c in walk(rep["Rows"]) if c[0].get("value") == f["row"])
    return float(cd[titles.index(f["column"])]["value"] or 0)


def _long_cell(f):
    """Value of one cell in a client's long-format workbook (read-only)."""
    wb = openpyxl.load_workbook(ROOT / f["file"], data_only=True, read_only=True)
    try:
        return wb[f["sheet"]][f["cell"]].value
    finally:
        wb.close()


def packet_for(client, month):
    """The review packet's numbers for a client month, computed in memory through the same source path as the pipeline."""
    key = (client, month)
    if key not in _packets:
        import monthly_review
        import run_client
        import sources
        cfg = run_client.load_config(client)
        ts = pd.Timestamp(f"{month}-01")
        am = sources.income_statement(cfg)
        monthly = abc.build_monthly_summary(am, cfg["totals"]).set_index("month")
        bs = sources.balance_sheet(cfg)
        cash = abc.cash_position(bs.loc[:, bs.columns <= ts], sources.cash_row(cfg)) if bs is not None else None
        _packets[key] = monthly_review.build_computed(am, monthly, ts, cash, categories=monthly_review.rollup_categories(am, cfg))
    return _packets[key]


def run_source_fields_case(case):
    """G34+: a packet number for a non-workbook-wide source (QuickBooks JSON or a long-format sheet) equals the value
    worked out by hand from the named source fields/cells, which are re-read here and must still hold the recorded value."""
    for f in case["source"].get("fields", []) + case["source"].get("cells", []):
        got = _report_field(f) if "row" in f else _long_cell(f)
        if not close(got, f["value"]):
            return False, f"source {f} now reads {got!r}"
    node = packet_for(case["input"]["client"], case["input"]["month"])
    for k in case["input"]["packet_path"].split("."):
        node = node[k]
    want = case["expected"]["value"]
    return close(node, want, case["expected"].get("tol", CENT)), f"packet {case['input']['packet_path']} = {node} (expected {want})"


def run_case(case):
    if case["type"] == "source_fields":
        return run_source_fields_case(case)
    if case["type"] == "yoy_view":
        return run_yoy_case(case)
    if case["type"] == "adjusted_net_income":
        return run_adjusted_case(case)
    if case["type"] == "commentary_sections":
        return run_sections_case(case)
    typ, data, expected = case["type"], case["input"], case["expected"]
    if typ == "intake_gate":
        return run_intake_case(case)
    if typ == "flux_item":
        return run_flux_item_case(case)
    if typ == "skill_text":
        return run_skill_text_case(case)
    if isinstance(case.get("source"), str) and case["source"].startswith("workbook"):
        return run_real_case(case)
    if typ.startswith("abc_"):
        return run_abc_case(typ, data, expected)
    if typ in {"revenue_mom", "zero_prior_revenue"}:
        df = pd.DataFrame({"month": data["months"], "revenue": data["revenue"]})
        out = normalize_monthly(df)
        actual_change = out.iloc[-1]["revenue_mom_change"]
        actual_pct = out.iloc[-1]["revenue_mom_pct"]
        ok = approx(actual_change, expected["absolute_change"])
        if expected["percentage_change"] is None:
            ok = ok and pd.isna(actual_pct)
        else:
            ok = ok and approx(actual_pct, expected["percentage_change"])
        return ok, f"change={actual_change}, pct={actual_pct}"
    if typ == "month_labels":
        out = normalize_monthly(pd.DataFrame({"month": data["months"], "revenue": data["revenue"]}))
        got = [f"{m:%Y-%m-%d}" for m in out["month"]]
        ok = got == expected["months"] and out["revenue"].tolist() == expected["revenue"]
        return ok, f"months={got}, revenue={out['revenue'].tolist()}"
    if typ == "gross_margin":
        out = normalize_monthly(pd.DataFrame([data]))
        actual = out.iloc[0]["gross_margin_pct"]
        return approx(actual, expected["gross_margin"]), f"gross_margin={actual}"
    if typ == "ttm":
        n = data["months_after"]
        months = pd.date_range(data["start_month"], periods=n, freq="MS")
        out = normalize_monthly(pd.DataFrame({"month": months, "revenue": [data["revenue_each_month"]] * n}))
        before = out.iloc[10]["revenue_ttm"]
        after = out.iloc[11]["revenue_ttm"]
        ok = pd.isna(before) and approx(after, expected["ttm_after_12"])
        return ok, f"ttm_after_11={before}, ttm_after_12={after}"
    if typ == "duplicate_month":
        try:
            normalize_monthly(pd.DataFrame({"month": data["months"], "revenue": data["revenue"]}))
        except ValueError as exc:
            return expected["raises"] in str(exc), str(exc)
        return False, "No validation error was raised"
    if typ == "missing_revenue":
        try:
            normalize_monthly(pd.DataFrame(data))
        except ValueError as exc:
            return expected["raises"] in str(exc), str(exc)
        return False, "No validation error was raised"
    if typ == "expense_category":
        df = pd.DataFrame({"month": data["months"], "expense_category": [data["category"]] * 2, "expense_amount": data["amounts"]})
        out = category_summary(df)
        last = out.iloc[-1]
        ok = approx(last["expense_mom_change"], expected["absolute_change"]) and approx(last["expense_mom_pct"], expected["percentage_change"])
        return ok, f"change={last['expense_mom_change']}, pct={last['expense_mom_pct']}"
    if typ == "short_history_commentary":
        months = pd.date_range("2025-01-01", periods=data["months"], freq="MS")
        out = normalize_monthly(pd.DataFrame({"month": months, "revenue": [data["revenue_each_month"]] * data["months"]}))
        text = " ".join(make_commentary(out))
        return expected["commentary_contains"].lower() in text.lower(), text
    return False, f"Unknown case type: {typ}"


def main():
    cases = json.loads(CASES.read_text(encoding="utf-8"))
    passed = 0
    print(f"Running {len(cases)} golden cases")
    for case in cases:
        try:
            ok, detail = run_case(case)
        except Exception as exc:  # Report the failure as an eval result, not a hidden crash.
            ok, detail = False, f"Unexpected {type(exc).__name__}: {exc}"
        print(f"{'PASS' if ok else 'FAIL'} {case['id']} — {case['name']}: {detail}")
        passed += int(ok)
    print(f"\nScore: {passed}/{len(cases)} ({passed / len(cases):.0%})")
    return 0 if passed == len(cases) else 1

if __name__ == "__main__":
    raise SystemExit(main())
