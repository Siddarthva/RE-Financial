"""Generate seven SYNTHETIC test clients, c04-c10 (fictional companies, made-up numbers). Regenerate:

    python clients/generate_synthetic.py

Each client gets clients/<id>/config.yaml and its own workbook. They vary on purpose: fiscal-year start, sheet and
row labels, month header format and order, wide or long layout (c08, c09), with or without a Balance Sheet, and one
client (c07) with a month missing from the history, which must be reported as a warning (declared in clients/batch.yaml).
Every workbook foots exactly (totals are the sum of the rounded lines) and, where there is a Balance Sheet, it balances
and Current Year Earnings equals year-to-date net income for the configured fiscal year.
"""
from __future__ import annotations

import random
from pathlib import Path

import pandas as pd
import yaml
from openpyxl import Workbook

HERE = Path(__file__).resolve().parent

CATS = {"Payroll & Benefits": ["Salaries", "Payroll Taxes"], "Occupancy": ["Rent"], "Technology": ["Software"],
        "Sales & Marketing": ["Advertising"]}
SPECS = {
    "c04": dict(name="Harbor Lane Bakery (synthetic)", fy=1, layout="wide", fmt="%b %Y", newest_first=True, end="2026-06", n=15,
                header="Account", sheet="Income Statement", sections=("Income", "Cost of Goods Sold", "Operating Expenses"),
                totals=("Total Income", "Total Cost of Goods Sold", "Gross Profit", "Total Operating Expenses", "Operating Income", "Net Income"),
                bs=True, scale=80_000),
    "c05": dict(name="Kestrel Analytics Ltd (synthetic)", fy=4, layout="wide", fmt="%b-%y", newest_first=False, end="2026-07", n=14,
                header="Line Item", sheet="P&L", sections=("Revenue", "Cost of Sales", "Expenses"),
                totals=("Total Revenue", "Total Cost of Sales", "Gross Margin", "Total Expenses", "Operating Profit", "Net Profit"),
                bs=False, scale=250_000),
    "c06": dict(name="Pinecrest Outfitters (synthetic)", fy=7, layout="wide", fmt="%B %Y", newest_first=False, end="2026-08", n=13,
                header="Particulars", sheet="Statement of Profit", sections=("Sales", "Direct Costs", "Overheads"),
                totals=("Total Sales", "Total Direct Costs", "Gross Result", "Total Overheads", "Operating Result", "Net Result"),
                bs=False, scale=40_000),
    "c07": dict(name="Mosaic Dental Group (synthetic)", fy=10, layout="wide", fmt="%m/%Y", newest_first=True, end="2026-06", n=15,
                header="Description", sheet="Monthly P&L", sections=("Fee Income", "Clinical Costs", "Practice Costs"),
                totals=("Total Fee Income", "Total Clinical Costs", "Contribution", "Total Practice Costs", "EBIT", "Profit"),
                bs=False, scale=120_000, drop_month="2025-11"),
    "c08": dict(name="Tidewater Logistics (synthetic)", fy=1, layout="long", fmt="%Y-%m", end="2026-06", n=15,
                columns={"account": "Account", "month": "Month", "amount": "Amount", "section": "Section"},
                sheet="Ledger", sections=("Income", "Cost of Goods Sold", "Operating Expenses"),
                totals=("Total Income", "Total Cost of Goods Sold", "Gross Profit", "Total Operating Expenses", "Operating Income", "Net Income"),
                bs=True, scale=300_000),
    "c09": dict(name="Quillfeather Publishing (synthetic)", fy=4, layout="long", fmt="%b %Y", end="2026-05", n=14,
                columns={"account": "GL Account", "month": "Period", "amount": "Value", "section": "Group"},
                sheet="Export", sections=("Revenue", "Direct Costs", "Overheads"),
                totals=("Revenue Total", "Direct Costs Total", "Gross Profit", "Overheads Total", "Operating Profit", "Net Profit"),
                bs=False, scale=60_000),
    "c10": dict(name="Vantage Robotics GmbH (synthetic)", fy=2, layout="wide", fmt="%Y-%m", newest_first=False, end="2026-07", n=14,
                header="Konto", sheet="GuV", sections=("Umsatzerloese", "Materialaufwand", "Betriebsaufwand"),
                totals=("Summe Umsatzerloese", "Summe Materialaufwand", "Rohertrag", "Summe Betriebsaufwand", "Betriebsergebnis", "Jahresergebnis"),
                bs=False, scale=500_000),
}
INCOME, COGS = ["Product Sales", "Service Revenue"], ["Materials", "Freight"]
OPEX = [a for accts in CATS.values() for a in accts]
INTEREST = "Interest Expense"


def months_for(spec):
    end = pd.Period(spec["end"], "M")
    ms = [end - k for k in range(spec["n"] - 1, -1, -1)]
    return [m for m in ms if str(m) != spec.get("drop_month")]


def numbers(cid, spec, months):
    rnd = random.Random(cid)
    s = spec["scale"]
    base = {"Product Sales": 0.7 * s, "Service Revenue": 0.3 * s, "Materials": 0.28 * s, "Freight": 0.05 * s,
            "Salaries": 0.25 * s, "Payroll Taxes": 0.03 * s, "Rent": 0.06 * s, "Software": 0.02 * s, "Advertising": 0.04 * s,
            INTEREST: 0.005 * s}
    lines = {}
    for a, b in base.items():
        lines[a] = [b if a == "Rent" else round(b * (1 + 0.015 * i) * rnd.uniform(0.9, 1.1), 2) for i in range(len(months))]
    t = lambda accts: [round(sum(lines[a][i] for a in accts), 2) for i in range(len(months))]
    inc, cogs, opex = t(INCOME), t(COGS), t(OPEX)
    gp = [round(x - y, 2) for x, y in zip(inc, cogs)]
    oi = [round(g - o, 2) for g, o in zip(gp, opex)]
    ni = [round(o - i, 2) for o, i in zip(oi, lines[INTEREST])]
    return lines, {"inc": inc, "cogs": cogs, "gp": gp, "opex": opex, "oi": oi, "ni": ni}


def statement_rows(spec, lines, tot):
    """[(label, section header or None, values)] in statement order; the same rows for wide and long."""
    s, T = spec["sections"], spec["totals"]
    rows = [(s[0], "HEADER", None)] + [(a, s[0], lines[a]) for a in INCOME] + [(T[0], None, tot["inc"])]
    rows += [(s[1], "HEADER", None)] + [(a, s[1], lines[a]) for a in COGS] + [(T[1], None, tot["cogs"])]
    rows += [(T[2], None, tot["gp"])]
    rows += [(s[2], "HEADER", None)] + [(a, s[2], lines[a]) for a in OPEX] + [(T[3], None, tot["opex"])]
    rows += [(T[4], None, tot["oi"]), (INTEREST, None, lines[INTEREST]), (T[5], None, tot["ni"])]
    return rows


def balance_sheet(spec, months, ni, rnd):
    fy = spec["fy"]
    ytd, run, cur = [], 0.0, None
    for m, v in zip(months, ni):
        year = m.year + (1 if m.month >= fy and fy != 1 else 0)
        if year != cur:
            cur, run = year, 0.0
        run = round(run + v, 2)
        ytd.append(run)
    cash = [round(spec["scale"] * rnd.uniform(1.5, 3.0), 2) for _ in months]
    recv = [round(spec["scale"] * rnd.uniform(0.4, 0.8), 2) for _ in months]
    re_ = [round(spec["scale"] * 0.9, 2)] * len(months)
    assets = [round(c + r, 2) for c, r in zip(cash, recv)]
    pay = [round(a - e - y, 2) for a, e, y in zip(assets, re_, ytd)]
    return [("Cash at Bank", cash), ("Total Cash and Cash Equivalents", cash), ("Receivables", recv), ("Total Assets", assets),
            ("Payables", pay), ("Retained Earnings", re_), ("Current Year Earnings", ytd),
            ("Total Liabilities and Equity", [round(p + e + y, 2) for p, e, y in zip(pay, re_, ytd)])]


def write_wide(path, spec, months, rows, bs_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = spec["sheet"]
    ws["A1"] = f"{spec['name']} - SYNTHETIC test data, fictional company"
    order = list(range(len(months)))[::-1] if spec["newest_first"] else list(range(len(months)))
    ws.append([])
    ws.append([spec["header"]] + [months[k].strftime(spec["fmt"]) for k in order])
    for label, sec, vals in rows:
        ws.append([label] + ([None] * len(order) if vals is None else [vals[k] for k in order]))
    if bs_rows:
        bs = wb.create_sheet("Balance Sheet")
        bs.append(["Account"] + [m.end_time.strftime("%d %b %Y") for m in months])
        for label, vals in bs_rows:
            bs.append([label] + vals)
    wb.save(path)


def write_long(path, spec, months, rows, bs_rows):
    cols = spec["columns"]
    wb = Workbook()
    ws = wb.active
    ws.title = spec["sheet"]
    ws.append([cols["account"], cols["month"], cols["amount"], cols["section"]])
    for label, sec, vals in rows:
        if vals is None:
            continue  # long format has no header rows; the section column carries the grouping
        for m, v in zip(months, vals):
            ws.append([label, m.strftime(spec["fmt"]), v, sec])
    if bs_rows:
        bs = wb.create_sheet("Balance Sheet")
        bs.append([cols["account"], cols["month"], cols["amount"], cols["section"]])
        for label, vals in bs_rows:
            for m, v in zip(months, vals):
                bs.append([label, m.strftime(spec["fmt"]), v, None])
    wb.save(path)


def config(cid, spec, book):
    s, T = spec["sections"], spec["totals"]
    layout = ({"type": "long", "month_format": spec["fmt"], "columns": spec["columns"]} if spec["layout"] == "long"
              else {"header_label": spec["header"], "month_format": spec["fmt"]})
    cfg = {"client": cid, "display_name": spec["name"], "workbook": book,
           "sheets": {"income_statement": spec["sheet"], **({"balance_sheet": "Balance Sheet"} if spec["bs"] else {})},
           "layout": layout,
           "sections": {"income": {"header": s[0], "total": T[0]}, "cost of goods sold": {"header": s[1], "total": T[1]},
                        "operating expenses": {"header": s[2], "total": T[2 + 1]}},
           "totals": dict(zip(["revenue", "cost_of_goods_sold", "gross_profit", "operating_expenses", "operating_income", "net_income"], T)),
           "subtotal_rows": [], "thresholds": {"gross_margin_drop_pp": 5},
           "fiscal_year": {"start_month": spec["fy"]}, "categories": list(CATS),
           "account_to_category": {a: c for c, accts in CATS.items() for a in accts}}
    head = (f"# {spec['name']} - SYNTHETIC test client (fictional company, made-up numbers).\n"
            f"# Generated by clients/generate_synthetic.py; do not edit by hand. Layout: {spec['layout']}.\n")
    return head + yaml.safe_dump(cfg, sort_keys=False, allow_unicode=True)


def main():
    for cid, spec in SPECS.items():
        months = months_for(spec)
        lines, tot = numbers(cid, spec, months)
        rows = statement_rows(spec, lines, tot)
        bs_rows = balance_sheet(spec, months, tot["ni"], random.Random(cid + "bs")) if spec["bs"] else None
        d = HERE / cid
        d.mkdir(exist_ok=True)
        book = f"{cid}_financials.xlsx"
        (write_long if spec["layout"] == "long" else write_wide)(d / book, spec, months, rows, bs_rows)
        (d / "config.yaml").write_text(config(cid, spec, book), encoding="utf-8")
        print(f"{cid}: {spec['layout']}, {len(months)} months to {spec['end']}, FY start {spec['fy']}"
              + (f", {spec['drop_month']} missing" if spec.get("drop_month") else ""))


if __name__ == "__main__":
    main()
