"""Synthetic LMN Traders workbook (fictional). Regenerate: python clients/lmn/generate_dataset.py"""
import random
from datetime import date
from pathlib import Path
import openpyxl

random.seed(7)
months = []
y, m = 2025, 8
for _ in range(13):  # Aug 2025 .. Aug 2026
    months.append(date(y, m, 1))
    m += 1
    if m > 12:
        y, m = y + 1, 1

sales = {"Wholesale Sales": 600000, "Retail Sales": 250000}
direct = {"Purchases": 420000, "Freight Inward": 38000}
overheads = {"Staff Costs": 130000, "Shop Rent": 45000, "Electricity": 9000,
             "Advertising": 22000, "Accounting Fees": 8000, "Vehicle Running": 14000}
fixed = {"Shop Rent"}

def val(base, i, name):
    if name in fixed:
        return float(base)
    return round(base * (1 + 0.02 * i) * random.uniform(0.92, 1.08), 2)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Statement of Profit"
ws["A1"] = "LMN Traders (fictional)"
ws["A2"] = "Monthly statement of profit, USD"
hdr = 4
ws.cell(hdr, 1, "Particulars")
for j, d in enumerate(months):
    ws.cell(hdr, 2 + j, d.strftime("%B %Y"))
r = hdr + 1
def block(title, items, total_label):
    global r
    ws.cell(r, 1, title); r += 1
    first = r
    for name, base in items.items():
        ws.cell(r, 1, name)
        for j in range(13):
            ws.cell(r, 2 + j, val(base, j, name))
        r += 1
    ws.cell(r, 1, total_label)
    rows = list(range(first, r))
    for j in range(13):
        ws.cell(r, 2 + j, round(sum(ws.cell(x, 2 + j).value for x in rows), 2))
    tot = r; r += 1
    return tot
def calc(label, fn):
    global r
    ws.cell(r, 1, label)
    for j in range(13):
        ws.cell(r, 2 + j, round(fn(2 + j), 2))
    rr = r; r += 1
    return rr
v = lambda row: (lambda c: ws.cell(row, c).value)
rev = block("Sales", sales, "Total Sales")
cogs = block("Direct Costs", direct, "Total Direct Costs")
gp = calc("Gross Result", lambda c: ws.cell(rev, c).value - ws.cell(cogs, c).value)
opx = block("Overheads", overheads, "Total Overheads")
oi = calc("Operating Result", lambda c: ws.cell(gp, c).value - ws.cell(opx, c).value)
ws.cell(r, 1, "Interest Paid")
for j in range(13):
    ws.cell(r, 2 + j, 5000.0)
intr = r; r += 1
calc("Net Result", lambda c: ws.cell(oi, c).value - ws.cell(intr, c).value)
ws.column_dimensions["A"].width = 28
out = Path(__file__).with_name("lmn_financials.xlsx")
wb.save(out)
print("wrote", out)
