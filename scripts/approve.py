"""Record a human Approve / Reject decision for one client period.

    python scripts/approve.py abc 2026-06 --reviewer "Name" --reason "Why" [--reject]

Writes outputs/<client>/approval_<period>.json and appends to outputs/<client>/audit_log.jsonl.
Rules: reviewer name and reason are required; approving needs the finance-reviewer verdict to be
"approve" and every question on the sign-off pack's Open Questions sheet to have a status (Open, Answered,
Closed ... anything but blank); once a period is approved it is locked and no further decision is accepted.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "outputs"
CLIENT = re.compile(r"^[a-z0-9][a-z0-9_-]{0,63}$")
PERIOD = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")


def fail(msg: str, code: int = 1) -> int:
    print(f"ERROR: {msg}")
    return code


def blank_status_questions(pack: Path) -> list[str]:
    """Numbers of the Open Questions rows whose Status cell is blank. Raises ValueError if the pack cannot be read."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(pack, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"cannot open the sign-off pack {pack.name} ({type(exc).__name__})") from None
    try:
        if "Open Questions" not in wb.sheetnames:
            return []  # drop_empty removes the sheet only when it holds nothing at all
        rows = list(wb["Open Questions"].iter_rows(values_only=True))
    finally:
        wb.close()
    head = [str(h).strip() if h is not None else "" for h in rows[0]] if rows else []
    if "Status" not in head or "Question" not in head:
        raise ValueError(f"{pack.name}: the Open Questions sheet has no Status or Question column")
    s, q = head.index("Status"), head.index("Question")
    return [str(r[0] if r[0] not in (None, "") else i) for i, r in enumerate(rows[1:], 1)
            if q < len(r) and r[q] not in (None, "", "None") and (s >= len(r) or r[s] is None or not str(r[s]).strip())]


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("client"), ap.add_argument("period")
    ap.add_argument("--reviewer", required=True), ap.add_argument("--reason", required=True)
    ap.add_argument("--reject", action="store_true")
    a = ap.parse_args()
    if not CLIENT.match(a.client) or not PERIOD.match(a.period):
        return fail("invalid client or period", 2)
    reviewer, reason = a.reviewer.strip(), a.reason.strip()
    if not reviewer or not reason:
        return fail("reviewer name and reason must not be empty", 2)
    d = OUT / a.client
    packet, review = d / f"review_packet_{a.period}.json", d / f"review_{a.period}.json"
    path = d / f"approval_{a.period}.json"
    if not packet.is_file():
        return fail(f"no review packet for {a.client} {a.period}; run the pipeline first")
    try:
        existing = json.loads(path.read_text(encoding="utf-8")) if path.is_file() else {}
        verdict = json.loads(review.read_text(encoding="utf-8")).get("verdict") if review.is_file() else None
    except ValueError as exc:  # never guess a decision from a damaged file
        return fail(f"cannot read an existing decision or verdict file as JSON ({exc.__class__.__name__}); "
                    f"check {path.name} and {review.name} in outputs/{a.client}/. Nothing was recorded.", 5)
    if existing.get("status") == "approved":
        return fail(f"{a.client} {a.period} is approved and locked", 3)
    status = "rejected" if a.reject else "approved"
    if status == "approved" and verdict != "approve":
        return fail(f"cannot approve: finance-reviewer verdict is {verdict!r}, not 'approve'", 4)
    if status == "approved":
        pack = d / f"monthly_review_{a.period}.xlsx"
        if not pack.is_file():
            return fail(f"cannot approve: no sign-off pack for {a.client} {a.period} ({pack.name}); build it first", 6)
        try:
            blank = blank_status_questions(pack)
        except ValueError as exc:
            return fail(f"cannot approve: {exc}", 6)
        if blank:
            return fail(f"cannot approve: {len(blank)} open question(s) have a blank status (question #{', #'.join(blank)}) in "
                        f"{pack.name}. Set each Status (for example Open, Answered or Closed) and run again.", 6)
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    record = {"client": a.client, "period": a.period, "status": status, "reviewer": reviewer, "reason": reason,
              "decided_at": now, "review_verdict": verdict,
              "packet_sha256": hashlib.sha256(packet.read_bytes()).hexdigest()}
    path.write_text(json.dumps(record, indent=2), encoding="utf-8")
    with open(d / "audit_log.jsonl", "a", encoding="utf-8") as f:
        f.write(json.dumps({"ts": now, "event": f"period_{status}", "client": a.client, "period": a.period,
                            "actor": reviewer, "reason": reason, "status": status}) + "\n")
    print(f"{a.client} {a.period}: {status} by {reviewer}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
