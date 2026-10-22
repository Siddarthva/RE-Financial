"""Where a client's numbers come from: a wide workbook, a long-format workbook, or QuickBooks report JSON.

Every path that produces numbers (run_client.run, monthly_review.prepare and the intake gate in close_artifacts)
reads through here, so all three sources reach the engine as the same account_month / balance-sheet frames.

    source: workbook (default)  layout.type: wide (default) -> abc_workbook_adapter.parse_wide_income_statement
    source: workbook            layout.type: long           -> abc_workbook_adapter.parse_long_income_statement
    source: qbo                                              -> qbo_adapter.load
"""
from __future__ import annotations

import pandas as pd

import abc_workbook_adapter as engine
import qbo_adapter

SOURCES = ("workbook", "qbo")
LAYOUT_TYPES = ("wide", "long")


def kind(cfg: dict) -> str:
    """'wide', 'long' or 'qbo'."""
    if cfg.get("source", "workbook") == "qbo":
        return "qbo"
    return (cfg.get("layout") or {}).get("type", "wide")


def layout(cfg: dict) -> dict:
    """The engine layout for the income statement (wide or long)."""
    lay = cfg["layout"]
    out = {"sheet": cfg["sheets"]["income_statement"], "header_label": lay.get("header_label"),
           "month_format": lay.get("month_format"),
           "sections": {sid: (v["header"], v["total"]) for sid, v in cfg["sections"].items()}}
    if kind(cfg) == "long":
        out["columns"] = lay["columns"]
        out["header_label"] = lay["columns"]["account"]
    return out


def bs_layout(cfg: dict) -> dict:
    if kind(cfg) == "long":
        return {**layout(cfg), "sheet": cfg["sheets"]["balance_sheet"], "sections": {}}
    return {**engine.ABC_BS_LAYOUT, "sheet": cfg["sheets"]["balance_sheet"]}


def cash_row(cfg: dict) -> str:
    return cfg.get("cash_row", engine.ABC_CASH_ROW)


def income_statement(cfg: dict) -> pd.DataFrame:
    k = kind(cfg)
    if k == "qbo":
        return qbo_adapter.load(cfg)[0]
    if k == "long":
        return engine.parse_long_income_statement(cfg["workbook_path"], layout(cfg))
    return engine.parse_wide_income_statement(cfg["workbook_path"], layout(cfg))


def balance_sheet(cfg: dict) -> pd.DataFrame | None:
    if not cfg["sheets"].get("balance_sheet"):
        return None
    k = kind(cfg)
    if k == "qbo":
        return qbo_adapter.load(cfg)[1]
    if k == "long":
        return engine.parse_long_income_statement(cfg["workbook_path"], bs_layout(cfg))
    return engine.parse_wide_income_statement(cfg["workbook_path"], bs_layout(cfg))


def locate_frame(frame: pd.DataFrame) -> dict:
    """A locator (as close_artifacts.locate returns) built from a parsed long/QuickBooks frame: each value's position
    is in attrs['cell_of'], so cell_ref() reads it from 'refs' instead of computing a column letter."""
    sections = frame.attrs.get("sections", {})
    return {"sheet": frame.attrs.get("sheet"), "header_row": frame.attrs.get("header_row", 1),
            "account_col": frame.attrs.get("account_col", 0),
            "month_cols": {f"{m:%Y-%m}": i for i, m in enumerate(frame.columns)},
            "accounts": {(a, sections.get(a)): r for a, r in frame.attrs.get("row_of", {}).items()},
            "refs": {(a, f"{m:%Y-%m}"): ref for (a, m), ref in frame.attrs.get("cell_of", {}).items()}}


def sheet_names(cfg: dict) -> list[str]:
    """Sheets (or QuickBooks report names) present in the source, for the intake card."""
    if kind(cfg) == "qbo":
        names = [cfg["sheets"]["income_statement"]]
        if cfg["sheets"].get("balance_sheet"):
            names.append(cfg["sheets"]["balance_sheet"])
        return names
    from openpyxl import load_workbook
    wb = load_workbook(cfg["workbook_path"], read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def frame_preflight(cfg: dict, month: str | None) -> list[str]:
    """Checks for long-format and QuickBooks sources once the file itself opened: parse both reports (any STOP the
    parser raises is returned as the problem), then the requested and prior month, the total rows and section totals."""
    try:
        am = income_statement(cfg)
        balance_sheet(cfg)
    except (ValueError, KeyError) as exc:
        return [str(exc)]
    sheet = cfg["sheets"]["income_statement"]
    problems, have = [], {f"{m:%Y-%m}" for m in am.columns}
    if month:
        prior = f"{pd.Period(month, 'M') - 1}"
        for m, role in ((month, "the requested month"), (prior, f"the prior month of {month}")):
            if m not in have:
                problems.append(f"{sheet}: no data for {m} ({role}); months found: {min(have)} to {max(have)}.")
    labels = {str(a).casefold() for a in am.index}
    for key, label in cfg["totals"].items():
        if label.casefold() not in labels:
            problems.append(f"{sheet}: total row {label!r} ({key}) not found. It is required; check for a renamed or deleted row.")
    for sid, v in cfg["sections"].items():
        if v["total"].casefold() not in labels:
            problems.append(f"{sheet}: section total row {v['total']!r} not found; the {sid} section cannot be closed or footed.")
    return problems
