"""Monthly analysis -> review packet, and review packet + commentary + reviewer verdict -> Excel sign-off pack.

    python src/monthly_review.py prepare <client> <YYYY-MM>
    python src/monthly_review.py build   <client> <YYYY-MM> --commentary FILE --review FILE

`prepare` runs the client's analysis and writes outputs/<client>/review_packet_<month>.json, the JSON that
commentary is checked against. `build` grades the commentary against that packet, merges the finance-reviewer's
verdict and writes outputs/<client>/monthly_review_<month>.xlsx. Both run inside the client's sandbox, so they
cannot read or write another client's folders.
"""
from __future__ import annotations
import argparse
import json
import re
import sys
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

SRC = Path(__file__).resolve().parent
sys.path.insert(0, str(SRC))
sys.path.insert(0, str(SRC.parent / "03_evals"))
import abc_workbook_adapter as engine  # noqa: E402
import close_artifacts  # noqa: E402
import excel_layout  # noqa: E402
import run_client  # noqa: E402
from client_sandbox import client_sandbox  # noqa: E402
from grade_commentary import grade  # noqa: E402

METRICS = ["revenue", "cost_of_goods_sold", "gross_profit", "operating_expenses", "operating_income", "net_income"]
MARGINS = {"gross_profit": "gross_margin_pct", "operating_income": "operating_margin_pct", "net_income": "net_margin_pct"}
MONTH = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
RUNWAY_WARN_MONTHS = 6


def _num(x):
    return None if pd.isna(x) else round(float(x), 6)


def build_computed(account_month, monthly, month, cash=None, gross_margin_drop_pp=None):
    """The JSON that commentary is graded against: every month's totals/MoM/TTM/margins and the chosen month's line items.

    `monthly` is the adapter summary indexed by month; `month` is a Timestamp that must have a prior month.
    """
    months = list(monthly.index)
    if month not in months or months.index(month) == 0:
        raise ValueError(f"{month:%Y-%m} needs a prior month in the data. Available: {months[0]:%Y-%m} (first, no prior) to {months[-1]:%Y-%m}.")
    prior = months[months.index(month) - 1]
    if prior != month - pd.DateOffset(months=1):
        raise ValueError(f"{month:%Y-%m} needs its prior month {month - pd.DateOffset(months=1):%Y-%m} in the data; "
                         f"the previous month present is {prior:%Y-%m}.")
    sections = account_month.attrs["sections"]
    out = {"latest_month": f"{month:%Y-%m}", "prior_month": f"{prior:%Y-%m}", "months": {}, "line_items": {}}
    for i, m in enumerate(months):
        rec = {}
        adjacent = i and months[i - 1] == m - pd.DateOffset(months=1)  # no MoM across a gap in the months
        for k in METRICS:
            v, p = monthly.at[m, k], (monthly.at[months[i - 1], k] if adjacent else None)
            rec[k] = {"value": _num(v), "ttm": _num(monthly.at[m, f"{k}_ttm"])}
            if p is not None:
                rec[k]["mom_change"] = _num(v - p)
                rec[k]["mom_pct"] = _num((v - p) / abs(p)) if p else None
        for k, col in MARGINS.items():
            rec[k]["margin"] = _num(monthly.at[m, col])
            if adjacent:
                rec[k]["margin_change_pp"] = _num((monthly.at[m, col] - monthly.at[months[i - 1], col]) * 100)
            if pd.notna(monthly.at[m, "revenue_ttm"]):
                rec[k]["ttm_margin"] = _num(monthly.at[m, f"{k}_ttm"] / monthly.at[m, "revenue_ttm"])
        out["months"][f"{m:%Y-%m}"] = rec
    for account in account_month.index:
        if not sections.get(account) or account.casefold().startswith("total "):
            continue
        v, p = account_month.at[account, month], account_month.at[account, prior]
        out["line_items"][account] = {"section": sections[account], "value": _num(v), "prior_value": _num(p),
                                      "mom_change": _num(v - p), "mom_pct": _num((v - p) / abs(p)) if p else None}
        if sections[account] == "operating expenses":
            out["line_items"][account]["share_of_opex"] = _num(v / monthly.at[month, "operating_expenses"])
    ranked = sorted(out["line_items"].items(), key=lambda kv: abs(kv[1]["mom_change"] or 0), reverse=True)
    out["top_movements"] = [{"account": a, "mom_change": d["mom_change"]} for a, d in ranked[:3]]
    if gross_margin_drop_pp is not None:
        out["gross_margin_drop_threshold_pp"] = gross_margin_drop_pp
        out["gross_margin_flags"] = engine.margin_drop_flags(monthly, gross_margin_drop_pp)
    if cash is not None:
        change = cash["cash"] - cash["prior_cash"]
        out["cash"] = {"month": f"{cash['month']:%Y-%m}", "value": _num(cash["cash"]), "prior_value": _num(cash["prior_cash"]),
                       "mom_change": _num(change), "mom_pct": _num(change / abs(cash["prior_cash"])),
                       "burn": {f"{n}m": _num(v) for n, v in cash["burn"].items()},
                       "runway_months": {f"{n}m": _num(v) for n, v in cash["runway_months"].items()}, "primary_window": "3m"}
    return out


def prepare(client, month_text, clients_root=run_client.CLIENTS_ROOT, outputs_root=run_client.OUTPUTS_ROOT):
    if not MONTH.match(month_text):
        raise ValueError(f"Month must look like 2026-06, got {month_text!r}")
    month = pd.Timestamp(f"{month_text}-01")
    with client_sandbox(client, clients_root, outputs_root):  # same intake gate as scripts/run_pipeline.py, before anything is written
        report, _, _ = close_artifacts.intake_gate(run_client.load_config(client, clients_root), month_text, engine)
        refused = close_artifacts.refusal(report, month_text)
        if refused:
            raise ValueError(refused)
    paths = run_client.run(client, clients_root, outputs_root)  # writes analysis.xlsx, summary.json, commentary.md
    with client_sandbox(client, clients_root, outputs_root):
        cfg = run_client.load_config(client, clients_root)
        layout = {"sheet": cfg["sheets"]["income_statement"], "header_label": cfg["layout"]["header_label"],
                  "month_format": cfg["layout"]["month_format"],
                  "sections": {sid: (v["header"], v["total"]) for sid, v in cfg["sections"].items()}}
        am = engine.parse_wide_income_statement(cfg["workbook_path"], layout)
        monthly = engine.build_monthly_summary(am, cfg["totals"])
        cash = None
        build_computed(am, monthly.set_index("month"), month)  # fail early, with a clear message, if the month has no prior month
        if cfg["sheets"].get("balance_sheet"):
            bs = engine.parse_wide_income_statement(cfg["workbook_path"], {**engine.ABC_BS_LAYOUT, "sheet": cfg["sheets"]["balance_sheet"]})
            cash = engine.cash_position(bs.loc[:, bs.columns <= month])
        packet = build_computed(am, monthly.set_index("month"), month, cash,
                                gross_margin_drop_pp=cfg.get("thresholds", {}).get("gross_margin_drop_pp"))
        packet["client"], packet["display_name"] = client, cfg["display_name"]
        packet["one_offs"] = [{"account": r.account, "section": r.section, "month": f"{r.month:%Y-%m}", "amount": r.amount}
                              for r in engine.find_one_off_items(am).itertuples()]
        packet["warnings"] = json.loads(paths["summary"].read_text(encoding="utf-8"))["warnings"]
        out = Path(outputs_root) / client / f"review_packet_{month_text}.json"
        out.write_text(json.dumps(packet, indent=2), encoding="utf-8")
    return out


STATEMENT = re.compile(r"^\s*(\d+)[.)]\s+(.*\S)\s*$")


def parse_statements(text):
    """Numbered lines ('1. ...') are the statements the reviewer and the sign-off column refer to."""
    out = [(int(m.group(1)), m.group(2)) for line in text.splitlines() if (m := STATEMENT.match(line))]
    if not out or [n for n, _ in out] != list(range(1, len(out) + 1)):
        raise ValueError("Commentary must contain statements numbered 1., 2., 3. ... in order, one per line.")
    return out


def auto_questions(packet):
    qs = []
    month = packet["latest_month"]
    for o in packet.get("one_offs", []):
        if o["month"] == month:
            qs.append(("Auto", f"{o['account']} is non-zero only in {month} ({o['amount']:,.2f}). Is it non-recurring, and is it in the right section?",
                       "One-off items distort trends, TTM and margins."))
    for w in packet.get("warnings", []):
        qs.append(("Auto (data quality)", w, "Reported by the analysis run."))
    cash = packet.get("cash")
    if cash:
        runway = cash["runway_months"].get(cash["primary_window"])
        if runway is not None and runway < RUNWAY_WARN_MONTHS:
            qs.append(("Auto (cash)", f"Runway is {runway:.1f} months at the {cash['primary_window']} average burn (cash {cash['value']:,.2f}). "
                       "Is there a funding plan or forecast to compare with?", "Runway is derived from past balance changes only; no forecast was supplied."))
    return qs


OK, BAD = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFC7CE")
NOTE_COLUMNS = {"Statement", "Reviewer note", "Issue", "Suggested fix", "Question", "Why it matters", "Answer"}
ID_FORMATS = {"#": excel_layout.PLAIN, "Statement": excel_layout.PLAIN}


def _sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(list(r))
    return excel_layout.polish_sheet(ws, wrap=NOTE_COLUMNS, formats=ID_FORMATS)


def _read(path, what, fix, as_json=False):
    """Read an input file, or raise ValueError saying which file and what to do next."""
    path = Path(path)
    if not path.is_file():
        raise ValueError(f"{what} not found: {path}. {fix}")
    text = path.read_text(encoding="utf-8")
    if not as_json:
        return text
    try:
        return json.loads(text)
    except ValueError as exc:
        raise ValueError(f"{what} is not valid JSON: {path} (line {getattr(exc, 'lineno', '?')}). {fix}") from None


def build(client, month_text, commentary_path, review_path, clients_root=run_client.CLIENTS_ROOT, outputs_root=run_client.OUTPUTS_ROOT):
    if not run_client.CLIENT_ID.match(client) or not MONTH.match(month_text):
        raise ValueError(f"Client must be an id like 'abc' and month like 2026-06; got {client!r} {month_text!r}.")
    with client_sandbox(client, clients_root, outputs_root):
        out_dir = Path(outputs_root) / client
        packet = _read(out_dir / f"review_packet_{month_text}.json", "Review packet",
                       f"Run 'python src/monthly_review.py prepare {client} {month_text}' first.", as_json=True)
        text = _read(commentary_path, "Commentary file", "Save the drafted commentary first (step 3 of /monthly-analysis).")
        review = _read(review_path, "Reviewer verdict", "Save the finance-reviewer's JSON verdict first (step 4 of /monthly-analysis).", as_json=True)
        statements = parse_statements(text)
        overall = grade(text, packet)
        notes = {n["statement"]: n for n in review.get("statement_notes", [])}
        rows, number_rows = [], []
        for n, line in statements:
            g = grade(line, packet)
            misses = [x["token"] for x in g["numbers"] if not x["matched"]]
            auto = "no numbers" if not g["numbers"] else ("all %d match" % len(g["numbers"]) if not misses else "MISMATCH: " + ", ".join(misses))
            note = notes.get(n, {})
            ok = not misses and note.get("status", "ok") == "ok"
            rows.append((n, line, auto, note.get("status", "not reviewed"), note.get("note", ""), "OK" if ok else "CHANGE", ""))
            number_rows += [(n, x["token"], x["kind"], "yes" if x["matched"] else "NO", x["fact"]) for x in g["numbers"]]
        auto_pass = overall["passed"]
        approved = review.get("verdict") == "approve" and auto_pass
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"
        facts = [("Client", packet["display_name"]), ("Month", packet["latest_month"]),
                 ("Verdict", "APPROVED" if approved else "CHANGES REQUIRED"),
                 ("Reviewer verdict", review.get("verdict", "missing")),
                 ("Automatic number check", f"{overall['numbers_matched']}/{overall['numbers_checked']} numbers match the JSON"),
                 ("Top 3 movements named", f"{sum(m['mentioned'] for m in overall['top_movements'])}/3"),
                 ("Cash risk covered", "yes" if overall["cash_risk"]["covered"] else "NO"),
                 ("Open changes", len(review.get("changes", []))),
                 ("Final sign-off (name / date)", ""), ("Source packet", f"review_packet_{month_text}.json")]
        ws.append(["Item", "Value"])
        for k, v in facts:
            ws.append([k, v])
        excel_layout.polish_sheet(ws, formats={"Value": excel_layout.PLAIN}, table=False)
        for (c,) in ws.iter_rows(min_row=2, max_col=1):
            c.font = Font(bold=True)
        ws["B4"].fill = OK if approved else BAD
        sheet = _sheet(wb, "Commentary Sign-off", ["#", "Statement", "Automatic number check", "Reviewer status", "Reviewer note", "Result", "Sign-off (name / date)"],
                       rows)
        for row in sheet.iter_rows(min_row=2):
            row[5].fill = OK if row[5].value == "OK" else BAD
        numbers = _sheet(wb, "Numbers Check", ["Statement", "Number as written", "Unit", "Matches JSON?", "Matched fact"], number_rows)
        for row in numbers.iter_rows(min_row=2):
            if row[2].value != "amount":  # percent / pp facts are stored in points (70.2), not fractions
                row[4].number_format = "0.0;[Red](0.0)"
        _sheet(wb, "Top Movements", ["Account", "MoM change", "Named in commentary?"],
               [(m["account"], m["mom_change"], "yes" if m["mentioned"] else "NO") for m in overall["top_movements"]])
        _sheet(wb, "Reviewer Changes", ["Statement", "Type", "Issue", "Suggested fix"],
               [(c.get("statement", ""), c.get("type", ""), c.get("issue", ""), c.get("fix", "")) for c in review.get("changes", [])] or [("", "", "None", "")])
        questions = [(q["source"] if isinstance(q, dict) and "source" in q else "Reviewer", q["question"], q.get("why", "")) for q in review.get("open_questions", [])]
        questions += auto_questions(packet)
        _sheet(wb, "Open Questions", ["#", "Source", "Question", "Why it matters", "Owner", "Answer", "Status"],
               [(i, s, q, w, "", "", "Open") for i, (s, q, w) in enumerate(questions, 1)] or [(1, "", "None", "", "", "", "")])
        excel_layout.drop_empty(wb)
        target = out_dir / f"monthly_review_{month_text}.xlsx"
        wb.save(target)
    return target, approved


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("prepare")
    p.add_argument("client"), p.add_argument("month")
    b = sub.add_parser("build")
    b.add_argument("client"), b.add_argument("month")
    b.add_argument("--commentary", required=True), b.add_argument("--review", required=True)
    args = ap.parse_args(argv)
    sys.stdout.reconfigure(encoding="utf-8")
    try:
        if args.cmd == "prepare":
            print(f"packet: {prepare(args.client, args.month)}")
        else:
            path, approved = build(args.client, args.month, args.commentary, args.review)
            print(f"workbook: {path}\nverdict: {'APPROVED' if approved else 'CHANGES REQUIRED'}")
    except (ValueError, FileNotFoundError, PermissionError) as exc:
        # Expected failures (refused intake, bad input, missing file, isolation) get one clear line, not a traceback.
        print(f"ERROR: {exc}")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
