"""Engine-side artifacts for the Close Review Workspace: intake checks, flux table, workbook locator.

All arithmetic lives here (and in the existing engine). The UI only reads the JSON these produce.
"""
from __future__ import annotations

import hashlib
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import sources

# Flux rule: a line item is flagged when it moves by at least FLUX_MIN_DOLLARS *and* FLUX_MIN_PCT
# (or has no % because the prior month was zero). Sign flips use a lower dollar floor.
FLUX_MIN_DOLLARS = 5000.0
FLUX_MIN_PCT = 0.10
FLUX_SIGN_FLIP_DOLLARS = 1000.0

REASONS = {
    "NEW_ITEM": ("high", "Prior month was zero; current month is not. Flagged when |$Δ| >= ${d:,.0f}."),
    "DROPPED_ITEM": ("high", "Current month is zero; prior month was not. Flagged when |$Δ| >= ${d:,.0f}."),
    "SIGN_FLIP": ("high", "Sign changed versus prior month. Flagged when |$Δ| >= ${f:,.0f}."),
    "LARGE_MOVE": ("medium", "|$Δ| >= ${d:,.0f} and |%Δ| >= {p:.0%}."),
}


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _layout(cfg: dict) -> dict:
    return sources.layout(cfg)


def locate(cfg: dict) -> dict:
    """Scan the income-statement sheet and return header/month/account cell positions (read-only)."""
    if sources.kind(cfg) != "wide":  # long-format and QuickBooks parsers record each value's position themselves
        return sources.locate_frame(sources.income_statement(cfg))
    lay = _layout(cfg)
    wb = load_workbook(cfg["workbook_path"], read_only=True, data_only=True)
    try:
        ws = wb[lay["sheet"]]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    header_row = next((i for i in range(min(15, len(rows))) if lay["header_label"] in [str(v) for v in rows[i]]), None)
    if header_row is None:
        return {"sheet": lay["sheet"], "header_row": None, "account_col": None, "month_cols": {}, "accounts": {}}
    account_col = [str(v) for v in rows[header_row]].index(lay["header_label"])
    month_cols = {}
    for j, v in enumerate(rows[header_row]):
        if j == account_col or v is None:
            continue
        p = pd.to_datetime(str(v), format=lay["month_format"], errors="coerce")
        if pd.notna(p):
            month_cols[f"{p:%Y-%m}"] = j
    headers = {h.casefold(): sid for sid, (h, _) in lay["sections"].items()}
    totals = {t.casefold(): sid for sid, (_, t) in lay["sections"].items()}
    accounts, section = {}, None
    for i in range(header_row + 1, len(rows)):
        name = rows[i][account_col] if account_col < len(rows[i]) else None
        if name is None or (isinstance(name, float) and pd.isna(name)):
            name = rows[i][0]
        if name is None:
            continue
        name = str(name).strip()
        if name.casefold() in headers:
            section = headers[name.casefold()]
            continue
        if section is not None and totals.get(name.casefold()) == section:
            section_for_row, section = section, None
        else:
            section_for_row = section
        accounts.setdefault((name, section_for_row), i + 1)
    return {"sheet": lay["sheet"], "header_row": header_row + 1, "account_col": account_col,
            "month_cols": month_cols, "accounts": accounts}


def cell_ref(loc: dict, account: str, section: str | None, month: str) -> str | None:
    if "refs" in loc:
        return loc["refs"].get((account, month))
    row = loc["accounts"].get((account, section)) or next(
        (r for (a, _), r in loc["accounts"].items() if a == account), None)
    col = loc["month_cols"].get(month)
    if row is None or col is None:
        return None
    return f"{loc['sheet']}!{get_column_letter(col + 1)}{row}"


def materiality(cfg: dict | None = None) -> dict:
    """The flux thresholds for one client: `thresholds.materiality` in its config, else the engine defaults.
    Shown in the flux JSON and the sign-off pack so the reviewer sees why each item was flagged."""
    own = ((cfg or {}).get("thresholds") or {}).get("materiality") or {}
    unknown = set(own) - {"min_dollars", "min_pct", "sign_flip_dollars"}
    if unknown:
        raise ValueError(f"Unknown materiality setting(s) in the client config: {sorted(unknown)}. "
                         "Use min_dollars, min_pct and sign_flip_dollars.")
    out = {"min_dollars": float(own.get("min_dollars", FLUX_MIN_DOLLARS)), "min_pct": float(own.get("min_pct", FLUX_MIN_PCT)),
           "sign_flip_dollars": float(own.get("sign_flip_dollars", FLUX_SIGN_FLIP_DOLLARS))}
    if min(out.values()) < 0:
        raise ValueError("Materiality thresholds must not be negative.")
    return {**out, "source": "client config" if own else "engine default"}


def build_flux(packet: dict, loc: dict, rule: dict | None = None) -> dict:
    month, prior = packet["latest_month"], packet["prior_month"]
    rule = rule or materiality()
    min_d, min_p, flip_d = rule["min_dollars"], rule["min_pct"], rule["sign_flip_dollars"]
    items = []
    for account, v in packet["line_items"].items():
        cur, pri = v.get("value") or 0.0, v.get("prior_value") or 0.0
        delta = v.get("mom_change")
        if delta is None:
            continue
        code = None
        if pri == 0 and cur != 0 and abs(delta) >= min_d:
            code = "NEW_ITEM"
        elif cur == 0 and pri != 0 and abs(delta) >= min_d:
            code = "DROPPED_ITEM"
        elif pri * cur < 0 and abs(delta) >= flip_d:
            code = "SIGN_FLIP"
        elif abs(delta) >= min_d and v.get("mom_pct") is not None and abs(v["mom_pct"]) >= min_p:
            code = "LARGE_MOVE"
        if not code:
            continue
        sev, rule_text = REASONS[code]
        items.append({
            "account": account, "section": v.get("section"), "reason_code": code, "severity": sev,
            "metric": f"{account} ({v.get('section')})", "current": cur, "prior": pri,
            "dollar_change": delta, "pct_change": v.get("mom_pct"),
            "source_cell_current": cell_ref(loc, account, v.get("section"), month),
            "source_cell_prior": cell_ref(loc, account, v.get("section"), prior),
            "threshold_rule": rule_text.format(d=min_d, f=flip_d, p=min_p),
        })
    items.sort(key=lambda r: abs(r["dollar_change"]), reverse=True)
    return {"client": packet.get("client"), "period": month, "prior_period": prior,
            "rule": {"min_dollars": min_d, "min_pct": min_p, "sign_flip_dollars": flip_d},
            "materiality": {**rule, "statement": (f"Flagged when the month-over-month change is at least ${min_d:,.0f} and {min_p:.0%}; "
                                                   f"a sign change is flagged from ${flip_d:,.0f}. Source: {rule['source']}.")},
            "items": items}


def _first(refs: list[str], n: int = 5) -> str:
    return ", ".join(refs[:n]) + (f" (+{len(refs) - n} more)" if len(refs) > n else "")


def preflight(cfg: dict, month: str | None = None) -> list[str]:
    """Checks that must pass before the workbook is parsed. Returns plain-language problems naming the file,
    sheet or cell; an empty list means the engine may read the workbook. Never modifies the file."""
    if sources.kind(cfg) == "qbo":  # report JSON: the adapter's shape checks are the preflight (STOP with a named cause)
        return sources.frame_preflight(cfg, month)
    path = Path(cfg["workbook_path"])
    try:
        with open(path, "rb") as f:
            head = f.read(4)
    except FileNotFoundError:
        return [f"Workbook not found: {path}. Check 'workbook' in the client's config.yaml."]
    except PermissionError:
        return [f"Workbook is open or locked by another program: {path}. Close it (for example in Excel) and re-run."]
    if head != b"PK\x03\x04":
        return [f"Not a valid .xlsx workbook (the file is not a zip package): {path.name}. Re-export it from the source system."]
    try:
        values = load_workbook(path, data_only=True)
        formulas = load_workbook(path, data_only=False)
    except Exception as exc:  # zip present but not a readable workbook
        return [f"Cannot open {path.name} as an .xlsx workbook ({type(exc).__name__}). Re-export it from the source system."]
    lay = _layout(cfg)
    wanted = [lay["sheet"]] + ([cfg["sheets"]["balance_sheet"]] if cfg["sheets"].get("balance_sheet") else [])
    missing = [s for s in wanted if s not in values.sheetnames]
    if missing:
        return [f"Sheet {s!r} not found in {path.name}; sheets present: {values.sheetnames}." for s in missing]
    problems = []
    for s in wanted:
        uncached = [c.coordinate for row in formulas[s].iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("=") and values[s][c.coordinate].value is None]
        if uncached:
            problems.append(f"{s}: {len(uncached)} formula cells have no saved value ({_first(uncached)}). "
                            "The file was saved without calculating; open it in Excel, save, and re-run.")
        errors = [f"{c.coordinate} {c.value}" for row in values[s].iter_rows() for c in row if c.data_type == "e"]
        if errors:  # pandas reads Excel error values as blanks, so they must be caught here
            problems.append(f"{s}: {len(errors)} cells hold an Excel error value ({_first(errors)}). Fix the formula and re-run.")
    if sources.kind(cfg) == "long":  # header, unknown columns, duplicate rows, months and totals: the long parser's checks
        values.close()
        formulas.close()
        return problems + sources.frame_preflight(cfg, month)
    ws = values[lay["sheet"]]
    rows = [list(r) for r in ws.iter_rows(max_row=15, values_only=True)]
    hr = next((i for i, r in enumerate(rows) if lay["header_label"] in [str(v) for v in r]), None)
    if hr is None:
        return problems + [f"{lay['sheet']}: header row with {lay['header_label']!r} not found in rows 1-15."]
    merged = [str(m) for m in ws.merged_cells.ranges if m.min_row <= hr + 1 <= m.max_row]
    if merged:
        problems.append(f"{lay['sheet']}: header cells are merged ({_first(merged)}); unmerge them so each month has its own header.")
    months: dict = {}
    for j, v in enumerate(rows[hr]):
        p = pd.to_datetime(str(v), format=lay["month_format"], errors="coerce") if v is not None else pd.NaT
        if pd.notna(p):
            months.setdefault(f"{p:%Y-%m}", []).append(f"{get_column_letter(j + 1)}{hr + 1}")
    for m, cells in months.items():
        if len(cells) > 1:
            problems.append(f"{lay['sheet']}: month {m} appears in more than one column ({', '.join(cells)}).")
    if month:
        prior = f"{pd.Period(month, 'M') - 1}"
        for m, role in ((month, "the requested month"), (prior, f"the prior month of {month}")):
            if m not in months:
                problems.append(f"{lay['sheet']}: no column for {m} ({role}) in header row {hr + 1}; "
                                f"months found: {min(months, default='-')} to {max(months, default='-')}.")
    labels = {str(r[c]).strip().casefold() for r in ws.iter_rows(min_row=hr + 2, values_only=True) for c in range(min(len(r), 3)) if r[c] is not None}
    for key, label in cfg["totals"].items():
        if label.casefold() not in labels:
            problems.append(f"{lay['sheet']}: total row {label!r} ({key}) not found. It is required; check for a renamed or deleted row.")
    for sid, (_, total) in lay["sections"].items():
        if total.casefold() not in labels:
            problems.append(f"{lay['sheet']}: section total row {total!r} not found; the {sid} section cannot be closed or footed.")
    values.close()
    formulas.close()
    return problems


def intake_gate(cfg: dict, month: str, engine) -> tuple[dict, pd.DataFrame, dict]:
    """The one intake gate for every path that produces numbers (run_pipeline and monthly_review.prepare).
    Raises ValueError if the file cannot be read safely; otherwise returns (report, account_month, locator).
    The caller must refuse to continue when refusal(report) is not None."""
    problems = preflight(cfg, month)
    if problems:
        raise ValueError("Intake refused the workbook: " + " | ".join(problems))
    am = sources.income_statement(cfg)
    monthly = engine.build_monthly_summary(am, cfg["totals"])
    loc = sources.locate_frame(am) if sources.kind(cfg) != "wide" else locate(cfg)
    return build_intake(cfg, loc, monthly, am), am, loc


def refusal(report: dict, month: str) -> str | None:
    if report["status"] == "pass":
        return None
    failed = [f"{c['check']} at {c['cell_ref']}" + (f" ({c['detail']})" if c["detail"] else "")
              for c in report["checks"] if not c["passed"]]
    return f"Intake checks failed (details in intake_{month}.json): " + " | ".join(failed)


DATA_CHECK_FIXES = (
    ("Income statement sheet present", "Restore the sheet name or set sheets.income_statement in the client config."),
    ("Header row found", "Check the header label in the client config against the workbook."),
    ("Month columns contiguous", "Remove or fill the gap between month columns so each month sits next to the previous one."),
    ("Total row present", "Restore the total row label or fix the matching name under totals in the client config."),
    ("Every amount is a number", "Replace the text or error in the listed cell with a number."),
    ("Gross profit", "Find the month where the subtotal was keyed or linked wrongly and correct it."),
    ("Operating income", "Find the month where the subtotal was keyed or linked wrongly and correct it."),
    ("foots to its lines", "A detail line or its total is wrong in the listed month; compare the lines with the total."),
    ("Balance Sheet rows present", "Restore the row label or set balance_sheet_rows in the client config."),
    ("Balance Sheet balances", "Find the month's line that was keyed or linked wrongly."),
    ("Current Year Earnings", "Check that the Balance Sheet and Income Statement come from the same export."),
)


def data_check_rows(report: dict) -> list[tuple]:
    """Rows for the pack's 'Data checks' sheet: check, PASS/FAIL, cell reference, fix hint (empty when it passed)."""
    rows = []
    for c in report["checks"]:
        hint = "" if c["passed"] else next((h for key, h in DATA_CHECK_FIXES if key in c["check"]), "Open the cell and compare it with the source export.")
        detail = c["detail"].split("Fix:")[0].strip(" ;") if c["detail"] else ""
        rows.append((c["check"], "PASS" if c["passed"] else "FAIL", c["cell_ref"], detail if not c["passed"] else "", hint))
    return rows


BS_ROWS = {"total_assets": "Total Assets", "total_liabilities_equity": "Total Liabilities and Equity",
           "current_year_earnings": "Current Year Earnings"}
BS_TOLERANCE = 0.01


def balance_sheet_checks(cfg: dict, monthly: pd.DataFrame, account_month: pd.DataFrame) -> list[dict]:
    """Balance Sheet checks for clients that have one: it balances every month, and Current Year Earnings equals
    year-to-date net income (fiscal year per config). Returns [] when the client has no Balance Sheet sheet.
    Each failure names the cells and what to fix, so the reviewer does not have to dig."""
    bs_sheet = cfg["sheets"].get("balance_sheet")
    if not bs_sheet:
        return []
    from abc_workbook_adapter import ABC_BS_LAYOUT
    rows = {**BS_ROWS, **cfg.get("balance_sheet_rows", {})}
    bs = sources.balance_sheet(cfg)
    if sources.kind(cfg) == "wide":
        bs_loc = locate({**cfg, "sections": {}, "sheets": {**cfg["sheets"], "income_statement": bs_sheet},
                         "layout": {"header_label": ABC_BS_LAYOUT["header_label"], "month_format": ABC_BS_LAYOUT["month_format"]}})
    else:
        bs_loc = sources.locate_frame(bs)
    out = []

    def add(name, ok, ref, detail=""):
        out.append({"check": name, "passed": bool(ok), "cell_ref": ref, "detail": detail})

    missing = [v for v in rows.values() if v not in bs.index]
    if missing:
        add("Balance Sheet rows present", False, bs_sheet,
            f"missing {', '.join(missing)}. Fix: restore the row label in the Balance Sheet or set balance_sheet_rows in the client config.")
        return out
    add("Balance Sheet rows present", True, bs_sheet, ", ".join(rows.values()))

    def ref(label, ts):
        return cell_ref(bs_loc, label, None, f"{ts:%Y-%m}") or f"{bs_sheet}!{label} {ts:%Y-%m}"

    diff = bs.loc[rows["total_assets"]] - bs.loc[rows["total_liabilities_equity"]]
    off = [(ts, d) for ts, d in diff.items() if not abs(d) <= BS_TOLERANCE]  # NaN counts as off
    add("Balance Sheet balances (assets = liabilities + equity, every month)", not off,
        _first([ref(rows["total_assets"], ts) for ts, _ in off]) if off else bs_sheet,
        "Fix: find the month's line that was keyed or linked wrongly; " + "; ".join(f"{ts:%Y-%m}: assets - liabilities & equity = {d:,.2f}" for ts, d in off[:3]) if off else "")

    ni = monthly.set_index("month")["net_income"]
    start = cfg["fiscal_year"]["start_month"]
    fy = ni.index.map(lambda m: m.year + (1 if m.month >= start and start != 1 else 0))
    ytd = ni.groupby(fy).cumsum()
    cye = bs.loc[rows["current_year_earnings"]]
    off = []
    for ts, v in ytd.items():
        got = cye.get(ts)
        if got is None or not abs(got - v) <= BS_TOLERANCE:
            off.append((ts, got, v))
    add("Current Year Earnings = year-to-date net income (every month)", not off,
        _first([ref(rows["current_year_earnings"], ts) for ts, _, _ in off]) if off else bs_sheet,
        "Fix: check that the Balance Sheet and Income Statement are from the same export; " + "; ".join(
            f"{ts:%Y-%m}: balance sheet {g if g is None else format(g, ',.2f')} vs YTD net income {v:,.2f}" for ts, g, v in off[:3]) if off else "")
    return out


def build_intake(cfg: dict, loc: dict, monthly: pd.DataFrame, account_month: pd.DataFrame) -> dict:
    """Workbook card data: sheets, period range, hash, and pass/fail checks with cell refs."""
    path = Path(cfg["workbook_path"])
    sheets = sources.sheet_names(cfg)
    checks = []

    def add(name, ok, ref, detail=""):
        checks.append({"check": name, "passed": bool(ok), "cell_ref": ref, "detail": detail})

    sheet = loc["sheet"]
    add("Income statement sheet present", sheet in sheets, sheet)
    hr = loc["header_row"]
    add("Header row found", hr is not None, f"{sheet}!{get_column_letter(loc['account_col'] + 1)}{hr}" if hr else sheet,
        sources.layout(cfg)["header_label"] or "")
    cols = sorted(loc["month_cols"].values())
    if "refs" in loc:  # long-format / QuickBooks: months are values, not column positions (a gap is a data-quality warning)
        add("Months read", len(cols) >= 2, sheet, f"{len(cols)} months ({min(loc['month_cols'])} to {max(loc['month_cols'])})")
    else:
        contiguous = bool(cols) and cols == list(range(cols[0], cols[-1] + 1))
        add("Month columns contiguous", contiguous,
            f"{sheet}!{get_column_letter(cols[0] + 1)}{hr}:{get_column_letter(cols[-1] + 1)}{hr}" if cols and hr else sheet,
            f"{len(cols)} month columns")
    latest = f"{monthly['month'].iloc[-1]:%Y-%m}"
    for key, label in cfg["totals"].items():
        ref = cell_ref(loc, label, None, latest)
        add(f"Total row present: {label}", ref is not None, ref or sheet, key)
    bad = account_month.attrs.get("bad_cells", [])
    add("Every amount is a number", not bad, _first([b[0] for b in bad]) if bad else sheet,
        "; ".join(f"{a} {m}: {why}" for _, a, m, why in bad[:3]))
    mon = monthly.set_index("month")
    for name, total, parts, sign in (("Gross profit = revenue - COGS", "gross_profit", ("revenue", "cost_of_goods_sold"), -1),
                                     ("Operating income = gross profit - opex", "operating_income", ("gross_profit", "operating_expenses"), -1)):
        diff = mon[parts[0]] + sign * mon[parts[1]] - mon[total]
        off = [f"{m:%Y-%m}" for m, d in diff.items() if not abs(d) <= 0.005]  # NaN counts as off
        add(f"{name} (every month)", not off, _first([cell_ref(loc, cfg["totals"][total], None, m) or m for m in off]) if off else sheet,
            f"off in {_first(off)}" if off else "")
    # Each section's detail lines must foot to its total row in every month; otherwise some number is wrong.
    sections = account_month.attrs.get("sections", {})
    sub = {s.casefold() for s in cfg.get("subtotal_rows", [])}
    for sid, v in cfg["sections"].items():
        detail = [a for a in account_month.index if sections.get(a) == sid and a.casefold() not in sub]
        # A wrong section header would otherwise leave that section's lines unassigned without any error.
        add(f"Section header matched: {v['header']}", bool(detail), sheet, f"{len(detail)} lines under it" if detail else f"no lines found under a row labelled {v['header']!r}")
        if v["total"] not in account_month.index or not detail:
            continue
        diff = account_month.loc[detail].sum(min_count=1) - account_month.loc[v["total"]]
        off = [(f"{m:%Y-%m}", d) for m, d in diff.items() if not abs(d) <= 0.01]
        add(f"{v['total']} foots to its lines (every month)", not off,
            _first([cell_ref(loc, v["total"], None, m) or m for m, _ in off]) if off else sheet,
            "; ".join(f"{m}: lines - total = {d:,.2f}" for m, d in off[:3]))
    checks += balance_sheet_checks(cfg, monthly, account_month)
    return {
        "workbook": path.name, "sha256": sha256_file(path), "sheets_detected": sheets,
        "period_start": f"{monthly['month'].iloc[0]:%Y-%m}", "period_end": latest, "months": len(monthly),
        "checks": checks, "status": "pass" if all(c["passed"] for c in checks) else "fail",
    }


def build_tables(cfg: dict, account_month: pd.DataFrame, engine) -> dict:
    """Per-month tables for the Financial Analysis tabs: revenue by stream and opex by category."""
    def clean(x):
        return None if pd.isna(x) else round(float(x), 2)

    months = [f"{m:%Y-%m}" for m in account_month.columns]
    sections = account_month.attrs.get("sections", {})
    streams = {a: {f"{m:%Y-%m}": clean(account_month.loc[a, m]) for m in account_month.columns}
               for a in account_month.index if sections.get(a) == "income"}
    trends = engine.build_expense_category_trends(account_month, cfg["subtotal_rows"])
    rollup = engine.build_category_rollup(trends, cfg["account_to_category"])
    cats: dict = {}
    for r in rollup.itertuples():
        cats.setdefault(r.category, {})[f"{r.month:%Y-%m}"] = clean(r.amount)
    return {"months": months, "revenue_by_stream": streams, "opex_by_category": cats}
